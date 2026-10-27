# import re, os, pyodbc, sys, openpyxl
# from datetime import datetime, timedelta
# import pandas as pd
# from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
# from openpyxl.utils import get_column_letter
 
# INPUT_FOLDER = r"D:\SANGKURIANG\MD_PROJECT\PromoResult\Input"
# OUTPUT_FOLDER = "./output"

# INPUT_COL_KODE = "KODE_PROMO"
# INPUT_COL_START = "Start"
# INPUT_COL_END = "End"
# INPUT_COL_JENIS = "JENIS_PROMO"
# PROMO_TYPE_COL = "Promo Type"
# PROMO_TYPE_LABEL = "ADD {pct}%"
# INPUT_EXTENSIONS = (".xlsx", ".xlsm", ".xls")

# # Kalau True, data promo (sheet RAW) juga dibatasi ke tanggal Start-End dari file input.
# FILTER_PROMO_QUERY_BY_DATE = True
# COUNTER_DEPT_PREFIXES = ("C", "G", "M", "S", "U")
 
# DB_DRIVER = "{ODBC Driver 17 for SQL Server}"
# DB_SERVER = "149.129.249.131"
# DB_DATABASE = "Y2_PROD"
# DB_USERNAME = "sa"
# DB_PASSWORD = "Eshoes09"
# USE_WINDOWS_AUTH = False
 
# FONT_NAME = "Arial"


# def build_query(kode_promo: str, promo_start=None, promo_end=None) -> str:
#     kode_sql = str(kode_promo).replace("'", "''")
#     date_filter = ""
#     if FILTER_PROMO_QUERY_BY_DATE and promo_start is not None and promo_end is not None:
#         date_filter = (
#             f"AND MLR_DATEPIECE >= '{promo_start:%Y%m%d} 00:00:00' "
#             f"AND MLR_DATEPIECE < '{(promo_end + timedelta(days=1)):%Y%m%d} 00:00:00'"
#         )
#     query = f"""
#             SELECT 
#             LIGNE.GL_NUMERO as [Invoice Number],
#             LIGNE.GL_QTEFACT as [Quantity],
#             LIGNE.GL_PUTTC as [Retail Price],
#             TYPEREMISE.GTR_LIBELLE as [Disc Description],
#             (MLR_TOTALBASE - MLR_MONTANTTTC) as [Selling Price],
#             MLR_DATEPIECE as [Document Data],  
#             LIGNE.GL_CODEARTICLE as [Item Code],
#             MLR_ETABLISSEMENT as [Store of the document], 
#             MLR_INDICEG as [Document Index], 
#             MLR_MONTANTHT as [Disc - Tax excl. amount], 
#             MLR_MONTANTHTDEV as [Disc - Tax excl. amount in currency], 
#             MLR_MONTANTTTC as [Disc - Tax incl. amount], 
#             MLR_MONTANTTTCDEV as [Disc - Tax incl. amount in currency], 
#             MLR_NATUREPIECEG as [Disc line type], 
#             MLR_NUMERO as [Doc number], 
#             MLR_NUMLIGNE as [Line number], 
#             MLR_NUMORDRE as [Order Number], 
#             MLR_ORGREMISE as [Origin of discount], 
#             MLR_RANG as [Ranking number], 
#             MLR_RANGLIE as [Line ranking number], 
#             MLR_REMISE as [Disc percentage], 
#             MLR_SOUCHE as [Disc line stub], 
#             MLR_TOTALBASE as [Disc basis], 
#             MLR_TOTALBASEDEV as [Disc basis price - currency], 
#             MLR_CODECOND as [Sales Condition],
#             LIGNE.GL_ARTICLE as [Item Identiier], 
#             LIGNE.GL_CAISSE as [Register of the line], 
#             LIGNE.GL_DEPOT as [Warehouse Doc], 
#             LIGNE.GL_ETABLISSEMENT as [Store - document line], 
#             LIGNE.GL_FAMILLENIV1 as [Brand], 
#             LIGNE.GL_FAMILLENIV2 as [Dept], 
#             LIGNE.GL_FAMILLENIV3 as [Segment],  
#             LIGNE.GL_FOURNISSEUR as [Countermark supplier], 
#             LIGNE.GL_LIBELLE as [Line Description], 
#             LIGNE.GL_REPRESENTANT as [Doc lin sales representative], 
#             LIGNE.GL_TIERS as [Third party of the line], 
#             LIGNE.GL_TYPEARTICLE as [Item type line], 
#             LIGNE.GL_TYPELIGNE as [Type of line], 
#             LIGNE.GL_STATUTENVOI as [Shipping status for items],
#             PIECE.GP_CONFIE as [Doc of consigned items],
#             ARTICLE.GA_CHARLIBRE1 as [Vendor Style Code], 
#             ARTICLE.GA_DESIGNATION1 as [Add item description], 
#             ARTICLE.GA_CHARLIBRE2 as [Dimension], 
#             ARTICLE.GA_CHARLIBRE3 as [Develop code], 
#             GA_CODEBARRE as [Item barcode], 
#             GA_CODEDIM1 as [Colour], 
#             GA_CODEDIM2 as [Size], 
#             GA_CODEDIM3 as [Dimension 3], 
#             GA_CODEDIM4 as [Dimension 4], 
#             GA_CODEDIM5 as [Dimension 5]
#             FROM 
#             LIGNEREMISE 
#             LEFT JOIN LIGNE ON (
#                 MLR_NATUREPIECEG = LIGNE.GL_NATUREPIECEG 
#                 AND MLR_SOUCHE = LIGNE.GL_SOUCHE 
#                 AND MLR_NUMERO = LIGNE.GL_NUMERO 
#                 AND MLR_INDICEG = LIGNE.GL_INDICEG 
#                 AND MLR_NUMLIGNE = LIGNE.GL_NUMLIGNE
#             ) 
#             LEFT JOIN PIECE ON GP_NATUREPIECEG = LIGNE.GL_NATUREPIECEG 
#             AND GP_SOUCHE = LIGNE.GL_SOUCHE 
#             AND GP_NUMERO = LIGNE.GL_NUMERO 
#             AND GP_INDICEG = LIGNE.GL_INDICEG 
#             LEFT JOIN ARTICLE ON GA_ARTICLE = LIGNE.GL_ARTICLE 
#             LEFT JOIN ARTICLECOMPL ON GA2_ARTICLE = LIGNE.GL_ARTICLE 
#             LEFT JOIN TIERS ON LIGNE.GL_TIERS = T_TIERS 
#             AND T_NATUREAUXI = 'CLI'
#             LEFT JOIN TYPEREMISE on GTR_TYPEREMISE = MLR_TYPEREMISE
#             WHERE 
#             MLR_NUMLIGNE > 0 
#             AND MLR_NUMORDRE > 0
#             AND MLR_CODECOND LIKE '%{kode_sql}%'
#             {date_filter}
#             ORDER BY MLR_DATEPIECE DESC
#     """
#     return query
 
# def run_query(query: str) -> pd.DataFrame:
#     if USE_WINDOWS_AUTH:
#         conn_str = (
#             f"DRIVER={DB_DRIVER};SERVER={DB_SERVER};DATABASE={DB_DATABASE};"
#             f"Trusted_Connection=yes;"
#         )
#     else:
#         conn_str = (
#             f"DRIVER={DB_DRIVER};SERVER={DB_SERVER};DATABASE={DB_DATABASE};"
#             f"UID={DB_USERNAME};PWD={DB_PASSWORD};"
#         )
 
#     print("[INFO] Menghubungkan ke database...")
#     conn = pyodbc.connect(conn_str)
#     try:
#         print("[INFO] Menjalankan query...")
#         df = pd.read_sql(query, conn)
#     finally:
#         conn.close()
 
#     print(f"[INFO] Query selesai. {len(df)} baris diterima.")
#     return df
 
# def build_store_sales_query(date_start, date_end, brand: str) -> str:
#     if not brand:
#         raise ValueError(
#             "Brand (GL_FAMILLENIV1) kosong - tidak bisa membangun query Store Sales."
#         )
 
#     brand_sql = "'" + str(brand).replace("'", "''") + "'"
#     date_start_str = f"{date_start:%Y%m%d} 00:00:00"
#     date_end_str = f"{(date_end + timedelta(days=1)):%Y%m%d} 00:00:00"
 
#     query = f"""
#             SELECT 
#             GL_ETABLISSEMENT as [Store Code], 
#             GL_QTEFACT as [Quantity], 
#             GL_TOTALTTC as [Total]
#             FROM 
#             GCLIGNEARTDIM 
#             LEFT OUTER JOIN UTILISAT US1 ON GL_CREATEUR = US1.US_UTILISATEUR 
#             WHERE 
#             (
#                 GP_TICKETANNULE <> 'X' 
#                 AND (
#                 GP_TYPECOMPTA IN ('TIC')
#                 ) 
#                 AND GL_DATEPIECE >= '{date_start_str}' 
#                 AND GL_DATEPIECE < '{date_end_str}' 
#                 AND (GL_NATUREPIECEG = 'FFO') 
#                 AND (
#                 (
#                     GL_TYPELIGNE <> 'COM' 
#                     AND GL_TYPELIGNE <> 'TOT'
#                 )
#                 ) 
#                 AND (
#                 GL_ETABLISSEMENT LIKE '[BCEGHLMRSTUVWZ]%'
#                 ) 
#                 AND (
#                 T_ETABLISSEMENT IN (
#                     '10000', 'B010', 'B020', 'B030', 'B040', 'B050', 'B060', 'B070', 'B080', 'B090', 
#                     'B100', 'B700', 'C060', 'CWH', 'E010', 'E040', 'E060', 'E070', 'E100', 'E105', 
#                     'E120', 'E130', 'E140', 'E141', 'E142', 'E145', 'E150', 'E155', 'E160', 'E170', 
#                     'E180', 'E182', 'E185', 'E190', 'E210', 'E220', 'E225', 'E240', 'E250', 'E252', 
#                     'E310', 'E315', 'E330', 'E410', 'E415', 'E420', 'E425', 'E450', 'E452', 'E460', 
#                     'E510', 'E520', 'E550', 'E552', 'E553', 'E560', 'E570', 'E610', 'E620', 'E630', 
#                     'E640', 'E650', 'E660', 'E670', 'E672', 'E680', 'E682', 'E685', 'E690', 'E710', 
#                     'E720', 'E730', 'E740', 'E810', 'E815', 'E820', 'E822', 'E850', 'E855', 'E880', 
#                     'E885', 'E932', 'E935', 'E940', 'E942', 'E950', 'EDE', 'G035', 'G060', 'H001', 
#                     'L100', 'M010', 'M020', 'M030', 'M080', 'M155', 'M190', 'M240', 'M310', 'M420', 
#                     'M515', 'M620', 'M730', 'N190', 'N430', 'N930', 'P620', 'R100', 'S010', 'S020', 
#                     'S030', 'S040', 'S100', 'S150', 'S185', 'S210', 'S220', 'S230', 'S330', 'S610', 
#                     'S630', 'S820', 'S930', 'T045', 'T130', 'T140', 'T150', 'T220', 'T310', 'T340', 
#                     'T420', 'T520', 'T560', 'T610', 'T660', 'T685', 'T710', 'T720', 'T730', 'T855', 
#                     'T880', 'T950', 'TEST', 'U002', 'U005', 'U195', 'V140', 'V182', 'V210', 'V220', 
#                     'V310', 'V415', 'V452', 'V460', 'V520', 'V550', 'V620', 'V630', 'V660', 'V710', 
#                     'V720', 'V730', 'V815', 'V820', 'V850', 'V880', 'V935', 'WMKR', 'X999', 'Z010', 
#                     'Z030'
#                 )
#                 ) 
#                 AND (T_NATUREAUXI = 'CLI') 
#                 AND (
#                 GL_FAMILLENIV1 IN ({brand_sql}) 
#                 AND GL_FAMILLENIV2 NOT IN ('PKG')
#                 )
#                 AND GL_QTEFACT <= 1
#             ) 
#             ORDER BY 
#             GP_DATEPIECE DESC, 
#             GL_ETABLISSEMENT, 
#             GP_REFINTERNE, 
#             GL_NUMLIGNE, 
#             GP_HEURECREATION DESC
#     """
#     return query
 
# def aggregate_sales_by_store(sales_df: pd.DataFrame) -> dict:
#     """
#     Jumlahkan Quantity & Total dari hasil build_store_sales_query(),
#     dikelompokkan per Store Code. Hasilnya dict
#     {store_code: {"qty": ..., "amount": ...}} yang dipakai untuk mengisi
#     kolom Total (periode promo) dan LW Sales per baris toko di SUMMARY.
#     """
#     if sales_df is None or sales_df.empty:
#         return {}
 
#     sales_df = sales_df.copy()
#     sales_df["Store Code"] = sales_df["Store Code"].astype(str).str.strip()
#     sales_df["Quantity"] = pd.to_numeric(sales_df["Quantity"], errors="coerce").fillna(0)
#     sales_df["Total"] = pd.to_numeric(sales_df["Total"], errors="coerce").fillna(0)
 
#     grouped = sales_df.groupby("Store Code", dropna=True)[["Quantity", "Total"]].sum()
#     return {
#         store: {"qty": float(row["Quantity"]), "amount": float(row["Total"])}
#         for store, row in grouped.iterrows()
#     }
 
# BULAN_ID = {
#     1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "Mei", 6: "Jun",
#     7: "Jul", 8: "Agu", 9: "Sep", 10: "Okt", 11: "Nov", 12: "Des",
# }
 
# HARI_ID = {
#     0: "Senin", 1: "Selasa", 2: "Rabu", 3: "Kamis",
#     4: "Jumat", 5: "Sabtu", 6: "Minggu",
# }
 
# # ------------------------------------------------------------
# # ATURAN PERIODE LW SALES
# # ------------------------------------------------------------
# #   1. Jarak dari tanggal akhir LW ke tanggal mulai promo minimal 1 minggu,
# #      dihitung INKLUSIF (kedua tanggal ikut dihitung).
# #      Contoh: akhir LW 13 Sep -> mulai promo 19 Sep = 13,14,...,19 = 7 hari (OK)
# #   2. Selisih periode LW = selisih periode promo (promo_end - promo_start),
# #      dan tanggal mulai LW jatuh di hari yang sama dengan mulai promo
# #      (mundur per kelipatan 7 hari).
# #   3. Periode LW tidak boleh cross dengan tanggal promo
# #      (otomatis terpenuhi oleh aturan 1).
# #   Dipilih jumlah minggu mundur PALING KECIL yang memenuhi semua aturan.
# #
# # Contoh:
# #   Promo 19 Sep      (selisih 0)  -> LW 12 Sep        (mundur 1 minggu)
# #   Promo 19-20 Sep   (selisih 1)  -> LW 12-13 Sep     (mundur 1 minggu)
# #   Promo 19-23 Sep   (selisih 4)  -> LW 5-9 Sep       (mundur 2 minggu)
# #   Promo 19-25 Sep   (selisih 6)  -> LW 5-11 Sep      (mundur 2 minggu)
# #   Promo 19-27 Sep   (selisih 8)  -> LW 5-13 Sep      (mundur 2 minggu)
# #   Promo 4-23 Agu    (selisih 19) -> LW 7-26 Jul      (mundur 4 minggu)
# LW_MIN_DISTANCE_DAYS = 7  # jarak minimal akhir LW -> mulai promo (inklusif)
 
# def lw_distance_days(lw_end, promo_start) -> int:
#     """Jarak inklusif dari akhir LW ke mulai promo (13 Sep -> 19 Sep = 7)."""
#     return (promo_start - lw_end).days + 1
 
# def compute_lw_weeks_back(duration_days: int) -> int:
#     """
#     Jumlah minggu mundur terkecil (minimal 1) supaya:
#         (promo_start - lw_end) + 1 >= LW_MIN_DISTANCE_DAYS
#     dengan lw_start = promo_start - 7*k dan lw_end = lw_start + duration_days.
#     Dari syarat: 7*k - duration_days + 1 >= LW_MIN_DISTANCE_DAYS
#               -> k >= (duration_days + LW_MIN_DISTANCE_DAYS - 1) / 7
#     """
#     needed = duration_days + LW_MIN_DISTANCE_DAYS - 1
#     return max(1, -(-needed // 7))  # ceil tanpa import math
 
# def compute_lw_period(promo_start, promo_end):
#     """Hitung periode LW dari tanggal Start & End promo (dari file input)."""
#     duration_days = (promo_end - promo_start).days  # selisih hari (4-23 Agu = 19)
#     weeks_back = compute_lw_weeks_back(duration_days)
 
#     lw_start = promo_start - timedelta(weeks=weeks_back)
#     lw_end = lw_start + timedelta(days=duration_days)
 
#     return lw_start, lw_end
 
# def format_date_range_label(prefix: str, start, end) -> str:
#     """
#     Format label rentang tanggal, mis.:
#       - beda bulan : "LW Sales (31 Jul - 7 Agu'25)"
#       - bulan sama : "LW Sales (3-9 Jun'26)"
#     """
#     yy = end.strftime("%y")
 
#     if (start.year, start.month) == (end.year, end.month):
#         return f"{prefix} ({start.day}-{end.day} {BULAN_ID[end.month]}'{yy})"
#     return (
#         f"{prefix} ({start.day} {BULAN_ID[start.month]} - "
#         f"{end.day} {BULAN_ID[end.month]}'{yy})"
#     )
 
# def compute_stock_date(promo_start) -> pd.Timestamp:
#     """
#     Tanggal Stock = 1 hari sebelum tanggal promo dimulai.
#     Contoh: promo mulai 7 Agustus 2026 -> Stock 6 Agustus 2026.
#     """
#     return promo_start - timedelta(days=1)
 
# def format_single_date_label(prefix: str, date) -> str:
#     """Format label satu tanggal, mis. "Stock 27 Jul 2026"."""
#     return f"{prefix} {date.day} {BULAN_ID[date.month]} {date.year}"
 
# # ============================================================
# # 2. BANGUN TEMPLATE PROMO RESULT (RAW + SUMMARY)
# # ============================================================
 
# HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
# SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
# TOTAL_FILL = PatternFill("solid", fgColor="FCE4D6")
# INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # kuning muda = sel input manual
# thin = Side(style="thin", color="BFBFBF")
# BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
 
# STOCK_HEADER_LABEL = "Stock"
# LW_SALES_HEADER_LABEL = "LW Sales"
 
# # Kalau True, toko yang punya penjualan brand di periode promo / LW tapi TIDAK
# # punya transaksi promo tetap dimunculkan di SUMMARY (kolom diskon = 0).
# # Default False: baris toko hanya dari data promo (seperti sebelumnya).
# INCLUDE_NON_PROMO_STORES = False
 
# def is_counter_dept_store(store_code) -> bool:
#     """True kalau kode toko diawali C, G, M, S, atau U (Counter / Dept Store)."""
#     return str(store_code).strip().upper().startswith(COUNTER_DEPT_PREFIXES)
 
 
# def sort_store_codes(store_codes) -> list:
#     """
#     Urutkan kode toko: Counter & Dept Store (C, G, M, S, U) dulu,
#     baru toko lainnya. Di dalam tiap grup diurutkan alfabetis.
#     """
#     return sorted(
#         store_codes,
#         key=lambda st: (0 if is_counter_dept_store(st) else 1, str(st).strip().upper()),
#     )
 
 
# # ============================================================
# # JENIS PROMO -> pecah "Special Offer" jadi ADD 10% / ADD 15% / dst.
# # ============================================================
 
# def parse_jenis_promo(text) -> list:
#     """
#     Pecah teks JENIS_PROMO per koma dan ambil persentase ADD-nya.
#       "BUY2 ADD10%,BUY3 OR MORE ADD15% ON ALL D30-D50"
#         -> [{"pct": 10.0, "label": "ADD 10%", "rule": "BUY2 ADD10%"},
#             {"pct": 15.0, "label": "ADD 15%", "rule": "BUY3 OR MORE ADD15% ON ALL D30-D50"}]
#     Urutan hasil = urutan di teks input (dipakai sebagai urutan kolom di SUMMARY).
#     """
#     if text is None or (isinstance(text, float) and pd.isna(text)):
#         return []
#     rules = []
#     for part in str(text).split(","):
#         part = part.strip()
#         m = re.search(r"ADD\s*(\d+(?:\.\d+)?)\s*%", part, flags=re.IGNORECASE)
#         if not m:
#             if part:
#                 print(f"[WARNING] JENIS_PROMO '{part}' tidak punya pola 'ADDxx%' - diabaikan.")
#             continue
#         pct = float(m.group(1))
#         if any(r["pct"] == pct for r in rules):
#             continue
#         rules.append({"pct": pct, "label": PROMO_TYPE_LABEL.format(pct=f"{pct:g}"), "rule": part})
#     return rules
 
 
# def apply_promo_type(df: pd.DataFrame, rules: list):
#     """
#     Tambah kolom PROMO_TYPE_COL ke df: baris yang Disc percentage-nya sama dengan
#     persentase ADD di JENIS_PROMO diberi label "ADD xx%". Baris yang tidak cocok
#     tetap pakai Disc Description aslinya (+ persentasenya) supaya tidak hilang.
#     Return: (df baru, urutan label untuk kolom SUMMARY)
#     """
#     df = df.copy()
#     if "Disc percentage" not in df.columns:
#         raise KeyError("Kolom 'Disc percentage' tidak ada di hasil query - JENIS_PROMO tidak bisa dipakai.")
#     pct = pd.to_numeric(df["Disc percentage"], errors="coerce").round(4)
#     label_by_pct = {round(r["pct"], 4): r["label"] for r in rules}
 
#     matched = pct.map(label_by_pct)
#     fallback = df["Disc Description"].astype(str) + pct.map(
#         lambda v: "" if pd.isna(v) else f" {v:g}%")
#     df[PROMO_TYPE_COL] = matched.fillna(fallback)
 
#     n_unmatched = int(matched.isna().sum())
#     if n_unmatched:
#         print(
#             f"[WARNING] {n_unmatched} baris Disc percentage-nya tidak ada di JENIS_PROMO "
#             f"({sorted(pct[matched.isna()].dropna().unique().tolist())}) - "
#             f"dilabeli dengan Disc Description aslinya."
#         )
#     for r in rules:
#         n = int((matched == r["label"]).sum())
#         print(f"[INFO]   {r['label']:<10} <- '{r['rule']}' : {n} baris")
 
#     order = [r["label"] for r in rules]
#     order += sorted(set(df[PROMO_TYPE_COL]) - set(order))
#     return df, order
 
 
# def build_promo_result_workbook(
#     df: pd.DataFrame,
#     output_path: str,
#     lw_sales_label: str = None,
#     stock_label: str = None,
#     lw_sales_map: dict = None,
#     total_sales_label: str = None,
#     total_sales_map: dict = None,
#     group_col: str = "Disc Description",
#     group_order: list = None,
#     jenis_promo_text: str = None,
# ):
#     # group_col  : kolom RAW untuk memecah kolom diskon di SUMMARY
#     #              ("Disc Description" default, atau PROMO_TYPE_COL kalau JENIS_PROMO diisi).
#     # group_order: urutan kolom diskon (default: alfabetis).
#     # Kalau label tidak dikirim (mis. dipanggil manual tanpa lewat main()),
#     # pakai fallback konstanta di bagian KONFIGURASI.
#     lw_sales_label = lw_sales_label or LW_SALES_HEADER_LABEL
#     stock_label = stock_label or STOCK_HEADER_LABEL
#     total_sales_label = total_sales_label or "Total"
#     lw_sales_map = lw_sales_map or {}
#     total_sales_map = total_sales_map or {}
 
#     required_cols = ["Store of the document", group_col, "Quantity", "Selling Price"]
#     for col in required_cols:
#         if col not in df.columns:
#             raise KeyError(f"Kolom '{col}' tidak ada di hasil query.")
 
#     df = df.dropna(subset=["Store of the document", group_col]).copy()
#     df["Quantity"] = df["Quantity"].fillna(0)
#     df["Selling Price"] = df["Selling Price"].fillna(0)
 
#     present = set(df[group_col].unique().tolist())
#     if group_order:
#         disc_set = [g for g in group_order if g in present]
#         disc_set += sorted(present - set(disc_set))
#     else:
#         disc_set = sorted(present)
#     store_set = sort_store_codes(df["Store of the document"].unique().tolist())
 
#     # Toko yang punya penjualan brand (periode promo / LW) tapi tidak ada di data promo
#     promo_stores_clean = {str(st).strip() for st in store_set}
#     extra_stores = sorted(
#         (set(total_sales_map) | set(lw_sales_map)) - promo_stores_clean
#     )
#     if extra_stores:
#         if INCLUDE_NON_PROMO_STORES:
#             print(f"[INFO] {len(extra_stores)} toko tanpa transaksi promo ikut ditampilkan: {extra_stores}")
#             store_set = sort_store_codes(set(store_set) | set(extra_stores))
#         else:
#             print(
#                 f"[WARNING] {len(extra_stores)} toko punya penjualan brand tapi tidak ada di data "
#                 f"promo, tidak ditampilkan di SUMMARY."
#                 # f" {extra_stores} (set INCLUDE_NON_PROMO_STORES = True untuk menampilkannya)."
#             )
 
#     out_wb = openpyxl.Workbook()
 
#     raw_ws = out_wb.active
#     raw_ws.title = "RAW"
#     all_cols = list(df.columns)
#     raw_ws.append(all_cols)
#     for cell in raw_ws[1]:
#         cell.font = Font(name=FONT_NAME, bold=True)
#     for row in df.itertuples(index=False, name=None):
#         raw_ws.append(list(row))
#     for c_idx in range(1, len(all_cols) + 1):
#         raw_ws.column_dimensions[get_column_letter(c_idx)].width = 20
#     raw_ws.freeze_panes = "A2"
 
#     # Cari posisi kolom (di sheet RAW) yang dibutuhkan formula SUMMARY,
#     # supaya tetap benar walau urutan/isi kolom RAW berubah.
#     store_col_letter = get_column_letter(all_cols.index("Store of the document") + 1)
#     disc_col_letter = get_column_letter(all_cols.index(group_col) + 1)
#     qty_col_letter = get_column_letter(all_cols.index("Quantity") + 1)
#     amt_col_letter = get_column_letter(all_cols.index("Selling Price") + 1)
 
#     raw_rows = len(df)
#     RAW_STORE_RNG = f"RAW!${store_col_letter}$2:${store_col_letter}${raw_rows + 1}"
#     RAW_DISC_RNG = f"RAW!${disc_col_letter}$2:${disc_col_letter}${raw_rows + 1}"
#     RAW_QTY_RNG = f"RAW!${qty_col_letter}$2:${qty_col_letter}${raw_rows + 1}"
#     RAW_AMT_RNG = f"RAW!${amt_col_letter}$2:${amt_col_letter}${raw_rows + 1}"
 
#     # --- SUMMARY sheet ---
#     sum_ws = out_wb.create_sheet("SUMMARY")
 
#     n_disc = len(disc_set)
#     # total_cols awal hanya dipakai untuk menulis judul; akan disesuaikan lagi
#     # (di-merge ulang) setelah kolom Stock/Sellthru/Cont%/LW Sales/Growth% dibuat.
#     sum_ws.cell(row=1, column=1, value="PROMO RESULT").font = Font(name=FONT_NAME, bold=True, size=14)
 
#     note = sum_ws.cell(
#         row=2, column=1,
#         value=(
#             f"Kolom diskon: Qty = SUM(Quantity), Amount = SUM(Selling Price) per Store x "
#             f"{group_col} (sheet RAW). Total & LW Sales: penjualan brand per toko selama periode "
#             "promo & periode LW (query sales toko)."
#         ),
#     )
#     note.font = Font(name=FONT_NAME, italic=True, size=9, color="808080")
#     if jenis_promo_text:
#         jp = sum_ws.cell(row=3, column=1, value=f"Jenis Promo: {jenis_promo_text}")
#         jp.font = Font(name=FONT_NAME, bold=True, size=10)
 
#     HEADER_ROW = 4
#     SUBHEADER_ROW = 5
#     DATA_START_ROW = 6
 
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=1, end_row=SUBHEADER_ROW, end_column=1)
#     c = sum_ws.cell(row=HEADER_ROW, column=1, value="Store Code")
#     c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     c.fill = HEADER_FILL
#     c.alignment = Alignment(horizontal="center", vertical="center")
#     c.border = BORDER
 
#     col = 2
#     disc_col_start = {}
#     for disc in disc_set:
#         sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
#         hc = sum_ws.cell(row=HEADER_ROW, column=col, value=disc)
#         hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#         hc.fill = HEADER_FILL
#         hc.alignment = Alignment(horizontal="center", vertical="center")
#         for cc in range(col, col + 2):
#             sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
 
#         qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=col, value="Qty")
#         amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=col + 1, value="Amount")
#         for cell in (qty_c, amt_c):
#             cell.font = Font(name=FONT_NAME, bold=True)
#             cell.fill = SUBHEADER_FILL
#             cell.alignment = Alignment(horizontal="center", vertical="center")
#             cell.border = BORDER
 
#         disc_col_start[disc] = col
#         col += 2
 
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
#     hc = sum_ws.cell(row=HEADER_ROW, column=col, value=total_sales_label)
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for cc in range(col, col + 2):
#         sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
#     total_qty_col = col
#     total_amt_col = col + 1
#     qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=total_qty_col, value="Qty")
#     amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=total_amt_col, value="Amount")
#     for cell in (qty_c, amt_c):
#         cell.font = Font(name=FONT_NAME, bold=True)
#         cell.fill = SUBHEADER_FILL
#         cell.alignment = Alignment(horizontal="center", vertical="center")
#         cell.border = BORDER
 
#     # --- Kolom tambahan: Stock | Sellthru | Cont% | LW Sales (Qty/Amount) | Growth% ---
#     col = total_amt_col + 1
 
#     stock_col = col
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=stock_col, end_row=SUBHEADER_ROW, end_column=stock_col)
#     hc = sum_ws.cell(row=HEADER_ROW, column=stock_col, value=stock_label)
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for rr in (HEADER_ROW, SUBHEADER_ROW):
#         sum_ws.cell(row=rr, column=stock_col).border = BORDER
#     col += 1
 
#     sellthru_col = col
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=sellthru_col, end_row=SUBHEADER_ROW, end_column=sellthru_col)
#     hc = sum_ws.cell(row=HEADER_ROW, column=sellthru_col, value="Sellthru")
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for rr in (HEADER_ROW, SUBHEADER_ROW):
#         sum_ws.cell(row=rr, column=sellthru_col).border = BORDER
#     col += 1
 
#     cont_col = col
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=cont_col, end_row=SUBHEADER_ROW, end_column=cont_col)
#     hc = sum_ws.cell(row=HEADER_ROW, column=cont_col, value="Cont%")
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for rr in (HEADER_ROW, SUBHEADER_ROW):
#         sum_ws.cell(row=rr, column=cont_col).border = BORDER
#     col += 1
 
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
#     hc = sum_ws.cell(row=HEADER_ROW, column=col, value=lw_sales_label)
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for cc in range(col, col + 2):
#         sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
#     lw_qty_col = col
#     lw_amt_col = col + 1
#     qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=lw_qty_col, value="Qty")
#     amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=lw_amt_col, value="Amount")
#     for cell in (qty_c, amt_c):
#         cell.font = Font(name=FONT_NAME, bold=True)
#         cell.fill = SUBHEADER_FILL
#         cell.alignment = Alignment(horizontal="center", vertical="center")
#         cell.border = BORDER
#     col += 2
 
#     growth_col = col
#     sum_ws.merge_cells(start_row=HEADER_ROW, start_column=growth_col, end_row=SUBHEADER_ROW, end_column=growth_col)
#     hc = sum_ws.cell(row=HEADER_ROW, column=growth_col, value="Growth%")
#     hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
#     hc.fill = HEADER_FILL
#     hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
#     for rr in (HEADER_ROW, SUBHEADER_ROW):
#         sum_ws.cell(row=rr, column=growth_col).border = BORDER
 
#     total_cols = growth_col  # update total_cols agar merge title/note di atas ikut menyesuaikan
#     sum_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
#     sum_ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
 
#     # baris grand total sudah diketahui sebelum loop, dipakai untuk formula Cont%
#     grand_row = DATA_START_ROW + len(store_set)
 
#     r = DATA_START_ROW
#     for store in store_set:
#         sc = sum_ws.cell(row=r, column=1, value=store)
#         sc.font = Font(name=FONT_NAME)
#         sc.alignment = Alignment(horizontal="center")
#         sc.border = BORDER
 
#         for disc in disc_set:
#             qcol = disc_col_start[disc]
#             acol = qcol + 1
#             qty_formula = (
#                 f'=SUMIFS({RAW_QTY_RNG},{RAW_STORE_RNG},$A{r},'
#                 f'{RAW_DISC_RNG},{get_column_letter(qcol)}$4)'
#             )
#             amt_formula = (
#                 f'=SUMIFS({RAW_AMT_RNG},{RAW_STORE_RNG},$A{r},'
#                 f'{RAW_DISC_RNG},{get_column_letter(qcol)}$4)'
#             )
#             qc = sum_ws.cell(row=r, column=qcol, value=qty_formula)
#             ac = sum_ws.cell(row=r, column=acol, value=amt_formula)
#             qc.number_format = '#,##0'
#             ac.number_format = '#,##0'
#             qc.font = Font(name=FONT_NAME)
#             ac.font = Font(name=FONT_NAME)
#             qc.border = BORDER
#             ac.border = BORDER
 
#         # --- Total Qty/Amount: penjualan brand toko ini selama periode promo,
#         #     dari query store sales (lihat aggregate_sales_by_store()).
#         #     Kalau toko tidak ketemu di hasil query, dianggap 0.
#         total_data = total_sales_map.get(str(store).strip(), {})
#         tq = sum_ws.cell(row=r, column=total_qty_col, value=total_data.get("qty", 0))
#         ta = sum_ws.cell(row=r, column=total_amt_col, value=total_data.get("amount", 0))
#         tq.number_format = '#,##0'
#         ta.number_format = '#,##0'
#         tq.font = Font(name=FONT_NAME, bold=True)
#         ta.font = Font(name=FONT_NAME, bold=True)
#         tq.border = BORDER
#         ta.border = BORDER
 
#         # --- Stock: sel input manual (kosong, isi nanti kalau sumber datanya sudah ada) ---
#         stock_c = sum_ws.cell(row=r, column=stock_col, value=None)
#         stock_c.fill = INPUT_FILL
#         stock_c.border = BORDER
#         stock_c.number_format = '#,##0'
 
#         # --- Sellthru = Total Qty / Stock (akan #DIV/0! selama Stock masih kosong) ---
#         tqty_ref = f"{get_column_letter(total_qty_col)}{r}"
#         stock_ref = f"{get_column_letter(stock_col)}{r}"
#         sellthru_c = sum_ws.cell(row=r, column=sellthru_col, value=f"={tqty_ref}/{stock_ref}")
#         sellthru_c.number_format = '0.0%'
#         sellthru_c.font = Font(name=FONT_NAME)
#         sellthru_c.border = BORDER
 
#         # --- Cont% = Total Qty toko ini / Grand Total Qty (lihat catatan asumsi di atas) ---
#         grand_qty_ref = f"${get_column_letter(total_qty_col)}${grand_row}"
#         cont_c = sum_ws.cell(row=r, column=cont_col, value=f"={tqty_ref}/{grand_qty_ref}")
#         cont_c.number_format = '0.0%'
#         cont_c.font = Font(name=FONT_NAME)
#         cont_c.border = BORDER
 
#         # --- LW Sales Qty/Amount: sekarang otomatis dari query LW Sales,
#         #     dijumlahkan per Store Code (lihat aggregate_sales_by_store()).
#         #     Kalau toko tidak ketemu di hasil query LW, dianggap 0.
#         lw_data = lw_sales_map.get(str(store).strip(), {})
#         lwq_c = sum_ws.cell(row=r, column=lw_qty_col, value=lw_data.get("qty", 0))
#         lwa_c = sum_ws.cell(row=r, column=lw_amt_col, value=lw_data.get("amount", 0))
#         for cell in (lwq_c, lwa_c):
#             cell.font = Font(name=FONT_NAME)
#             cell.border = BORDER
#             cell.number_format = '#,##0'
 
#         # --- Growth% = Total Amount / LW Amount - 1 (akan #DIV/0! selama LW Amount kosong) ---
#         tamt_ref = f"{get_column_letter(total_amt_col)}{r}"
#         lwamt_ref = f"{get_column_letter(lw_amt_col)}{r}"
#         growth_c = sum_ws.cell(row=r, column=growth_col, value=f"={tamt_ref}/{lwamt_ref}-1")
#         growth_c.number_format = '0.0%'
#         growth_c.font = Font(name=FONT_NAME)
#         growth_c.border = BORDER
 
#         r += 1
 
#     last_data_row = r - 1
#     assert grand_row == r, "grand_row seharusnya sama dengan baris setelah loop store"
 
#     gt = sum_ws.cell(row=r, column=1, value="GRAND TOTAL")
#     gt.font = Font(name=FONT_NAME, bold=True)
#     gt.fill = TOTAL_FILL
#     gt.border = BORDER
#     for cc in range(2, total_qty_col):  # sebelum kolom Total: SUM per kolom Qty/Amount disc
#         col_letter = get_column_letter(cc)
#         cell = sum_ws.cell(row=r, column=cc, value=f"=SUM({col_letter}{DATA_START_ROW}:{col_letter}{last_data_row})")
#         cell.number_format = '#,##0'
#         cell.font = Font(name=FONT_NAME, bold=True)
#         cell.fill = TOTAL_FILL
#         cell.border = BORDER
#     for cc in (total_qty_col, total_amt_col, stock_col, lw_qty_col, lw_amt_col):
#         col_letter = get_column_letter(cc)
#         cell = sum_ws.cell(row=r, column=cc, value=f"=SUM({col_letter}{DATA_START_ROW}:{col_letter}{last_data_row})")
#         cell.number_format = '#,##0'
#         cell.font = Font(name=FONT_NAME, bold=True)
#         cell.fill = TOTAL_FILL
#         cell.border = BORDER
#     # Sellthru / Cont% / Growth% di baris Grand Total dihitung ulang dari total-nya (bukan SUM dari %)
#     gt_tqty_ref = f"{get_column_letter(total_qty_col)}{r}"
#     gt_tamt_ref = f"{get_column_letter(total_amt_col)}{r}"
#     gt_stock_ref = f"{get_column_letter(stock_col)}{r}"
#     gt_lwamt_ref = f"{get_column_letter(lw_amt_col)}{r}"
#     for cc, formula in (
#         (sellthru_col, f"={gt_tqty_ref}/{gt_stock_ref}"),
#         (cont_col, "=1"),
#         (growth_col, f"={gt_tamt_ref}/{gt_lwamt_ref}-1"),
#     ):
#         cell = sum_ws.cell(row=r, column=cc, value=formula)
#         cell.number_format = '0.0%'
#         cell.font = Font(name=FONT_NAME, bold=True)
#         cell.fill = TOTAL_FILL
#         cell.border = BORDER
#         cell.border = BORDER
 
#     sum_ws.column_dimensions["A"].width = 14
#     for cc in range(2, total_cols + 1):
#         sum_ws.column_dimensions[get_column_letter(cc)].width = 16
#     sum_ws.freeze_panes = "B6"
 
#     out_wb.save(output_path)
#     print(f"[INFO] File tersimpan: {output_path}")
#     print(f"[INFO] Disc types: {disc_set}")
#     print(f"[INFO] Jumlah store: {len(store_set)}")
 
# # ============================================================
# # 3. BACA FILE INPUT
# # ============================================================
 
# def _norm(text) -> str:
#     return str(text).strip().lower().replace(" ", "").replace("_", "")
 
 
# def _clean_kode_promo(value):
#     """6988 / 6988.0 / ' 6988 ' -> '6988'. Kosong -> None."""
#     if value is None or (isinstance(value, float) and pd.isna(value)):
#         return None
#     if isinstance(value, float) and value.is_integer():
#         value = int(value)
#     kode = str(value).strip()
#     if kode.endswith(".0") and kode[:-2].isdigit():
#         kode = kode[:-2]
#     return kode or None
 
 
# def _parse_date(value):
#     if value is None or (isinstance(value, float) and pd.isna(value)):
#         return None
#     if isinstance(value, (datetime, pd.Timestamp)):
#         return pd.Timestamp(value).normalize()
#     ts = pd.to_datetime(str(value).strip(), dayfirst=True, errors="coerce")
#     return None if pd.isna(ts) else ts.normalize()
 
 
# def read_input_file(path: str) -> list:
#     """
#     Baca 1 file input. Baris 1 = header (KODE_PROMO | Start | End | JENIS_PROMO opsional),
#     data promo dimulai dari baris setelah header.
#     Return: list of dict {kode, start, end, source}.
#     """
#     raw = pd.read_excel(path, sheet_name=0, dtype=object)
#     if raw.empty:
#         return []
 
#     cols = {_norm(c): c for c in raw.columns}
#     try:
#         col_kode = cols[_norm(INPUT_COL_KODE)]
#         col_start = cols[_norm(INPUT_COL_START)]
#         col_end = cols[_norm(INPUT_COL_END)]
#     except KeyError:
#         if raw.shape[1] < 3:
#             raise ValueError(
#                 f"File {os.path.basename(path)} harus punya minimal 3 kolom "
#                 f"({INPUT_COL_KODE}, {INPUT_COL_START}, {INPUT_COL_END})."
#             )
#         print(
#             f"[WARNING] Header '{INPUT_COL_KODE}/{INPUT_COL_START}/{INPUT_COL_END}' tidak "
#             f"lengkap di {os.path.basename(path)} - dipakai kolom A, B, C."
#         )
#         col_kode, col_start, col_end = raw.columns[:3]
 
#     col_jenis = cols.get(_norm(INPUT_COL_JENIS))  # kolom opsional
 
#     promos = []
#     for idx, row in raw.iterrows():
#         excel_row = idx + 2  # +1 header, +1 karena index mulai 0
#         kode = _clean_kode_promo(row[col_kode])
#         start = _parse_date(row[col_start])
#         end = _parse_date(row[col_end])
 
#         if kode is None and start is None and end is None:
#             continue  # baris kosong
#         if kode is None or start is None or end is None:
#             print(
#                 f"[WARNING] {os.path.basename(path)} baris {excel_row}: data tidak lengkap / "
#                 f"tanggal tidak valid (kode={row[col_kode]}, start={row[col_start]}, "
#                 f"end={row[col_end]}). Dilewati."
#             )
#             continue
#         if end < start:
#             print(
#                 f"[WARNING] {os.path.basename(path)} baris {excel_row}: tanggal End "
#                 f"({end:%d %b %Y}) sebelum Start ({start:%d %b %Y}). Dilewati."
#             )
#             continue
 
#         jenis = None
#         if col_jenis is not None and not pd.isna(row[col_jenis]):
#             jenis = str(row[col_jenis]).strip() or None
 
#         promos.append({"kode": kode, "start": start, "end": end, "jenis": jenis,
#                        "source": f"{os.path.basename(path)} baris {excel_row}"})
#     return promos
 
 
# def read_all_inputs(folder: str) -> list:
#     if not os.path.isdir(folder):
#         raise FileNotFoundError(f"Folder input tidak ditemukan: {folder}")
 
#     files = sorted(
#         f for f in os.listdir(folder)
#         if f.lower().endswith(INPUT_EXTENSIONS) and not f.startswith("~$")  # skip file lock Excel
#     )
#     if not files:
#         print(f"[WARNING] Tidak ada file Excel di folder input: {folder}")
#         return []
 
#     promos = []
#     for f in files:
#         path = os.path.join(folder, f)
#         print(f"[INFO] Membaca file input: {path}")
#         try:
#             found = read_input_file(path)
#         except Exception as e:
#             print(f"[ERROR] Gagal membaca {f}: {e}")
#             continue
#         print(f"[INFO]   -> {len(found)} promo ditemukan.")
#         promos.extend(found)
#     return promos
 
 
# # ============================================================
# # 4. PROSES 1 PROMO
# # ============================================================
 
# def process_promo(kode_promo: str, promo_start, promo_end, jenis_promo: str = None):
#     print("=" * 70)
#     print(f"[INFO] KODE PROMO: {kode_promo} | {promo_start:%d %b %Y} - {promo_end:%d %b %Y}")
#     print("=" * 70)
 
#     # 1. Data promo (sheet RAW)
#     df = run_query(build_query(kode_promo, promo_start, promo_end))
#     if df.empty:
#         print(f"[WARNING] Tidak ada data promo untuk kode {kode_promo}. Dilewati.")
#         return None
 
#     # 1b. Pecah diskon berdasarkan JENIS_PROMO (kalau diisi di file input)
#     group_col, group_order = "Disc Description", None
#     rules = parse_jenis_promo(jenis_promo)
#     if rules:
#         print(f"[INFO] JENIS_PROMO: {jenis_promo}")
#         df, group_order = apply_promo_type(df, rules)
#         group_col = PROMO_TYPE_COL
#     elif jenis_promo:
#         print("[WARNING] JENIS_PROMO diisi tapi tidak ada 'ADDxx%' yang terbaca - pakai Disc Description.")
 
#     # 2. Periode Total / LW / Stock - semuanya dari tanggal di file input
#     duration_days = (promo_end - promo_start).days
#     lw_start, lw_end = compute_lw_period(promo_start, promo_end)
#     lw_sales_label = format_date_range_label("LW Sales", lw_start, lw_end)
#     total_sales_label = format_date_range_label("Total", promo_start, promo_end)
 
#     stock_date = compute_stock_date(promo_start)
#     stock_label = format_single_date_label("Stock", stock_date)
 
#     print(
#         f"[INFO] Periode promo (input)   : {promo_start:%d %b %Y} - {promo_end:%d %b %Y} "
#         f"({duration_days} hari, mulai hari {HARI_ID[promo_start.weekday()]})"
#     )
#     print(
#         f"[INFO] Periode LW Sales        : {lw_start:%d %b %Y} - {lw_end:%d %b %Y} "
#         f"(mundur {compute_lw_weeks_back(duration_days)} minggu, "
#         f"jarak {lw_distance_days(lw_end, promo_start)} hari ke mulai promo)"
#     )
#     print(f"[INFO] Tanggal Stock (H-1 promo): {stock_date:%d %b %Y}")
 
#     # 3. Sales toko periode promo & LW
#     brands_found = sorted(df["Brand"].dropna().astype(str).str.strip().unique().tolist())
#     if not brands_found:
#         print("[WARNING] Kolom 'Brand' kosong di data promo - Total & LW Sales dianggap 0 semua.")
#         total_sales_map, lw_sales_map = {}, {}
#     else:
#         if len(brands_found) > 1:
#             print(
#                 f"[WARNING] Ditemukan {len(brands_found)} Brand berbeda: {brands_found}. "
#                 f"Dipakai brand pertama saja: '{brands_found[0]}'."
#             )
#         brand = brands_found[0]
#         print(f"[INFO] Brand untuk query Store Sales (GL_FAMILLENIV1): '{brand}'")
 
#         print("[INFO] Query sales periode PROMO (kolom Total)...")
#         total_df = run_query(build_store_sales_query(promo_start, promo_end, brand))
#         total_sales_map = aggregate_sales_by_store(total_df)
#         print(f"[INFO] Data sales periode promo: {len(total_df)} baris, {len(total_sales_map)} store.")
 
#         print("[INFO] Query sales periode LW (kolom LW Sales)...")
#         lw_df = run_query(build_store_sales_query(lw_start, lw_end, brand))
#         lw_sales_map = aggregate_sales_by_store(lw_df)
#         print(f"[INFO] Data sales periode LW: {len(lw_df)} baris, {len(lw_sales_map)} store.")
 
#     # 4. Bangun file Promo Result
#     output_filename = f"Promo_Result_{kode_promo}_{datetime.now():%d_%b_%Y}.xlsx"
#     output_path = os.path.join(OUTPUT_FOLDER, output_filename)
#     build_promo_result_workbook(
#         df,
#         output_path,
#         lw_sales_label=lw_sales_label,
#         stock_label=stock_label,
#         lw_sales_map=lw_sales_map,
#         total_sales_label=total_sales_label,
#         total_sales_map=total_sales_map,
#         group_col=group_col,
#         group_order=group_order,
#         jenis_promo_text=jenis_promo,
#     )
#     return output_path
 
 
# def main():
#     os.makedirs(OUTPUT_FOLDER, exist_ok=True)
 
#     promos = read_all_inputs(INPUT_FOLDER)
#     if not promos:
#         print("[WARNING] Tidak ada promo valid di file input. Proses dihentikan.")
#         sys.exit(0)
 
#     print(f"[INFO] Total promo yang akan diproses: {len(promos)}")
 
#     success, failed = [], []
#     for p in promos:
#         try:
#             out = process_promo(p["kode"], p["start"], p["end"], p.get("jenis"))
#             if out:
#                 success.append((p["kode"], out))
#             else:
#                 failed.append((p["kode"], "tidak ada data promo"))
#         except Exception as e:
#             print(f"[ERROR] Kode promo {p['kode']} ({p['source']}) gagal: {e}")
#             failed.append((p["kode"], str(e)))
 
#     print("=" * 70)
#     print(f"[INFO] SELESAI. Berhasil: {len(success)} | Gagal/dilewati: {len(failed)}")
#     for kode, out in success:
#         print(f"   [OK]   {kode} -> {out}")
#     for kode, reason in failed:
#         print(f"   [SKIP] {kode} -> {reason}")
 
 
# if __name__ == "__main__":
#     main()


import os, re, pyodbc, sys, openpyxl
from datetime import datetime, timedelta
import pandas as pd
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
 
INPUT_FOLDER = r"D:\EVERBEST\NEW_GENERATOR\webapp\promo_result\Input"
OUTPUT_FOLDER = "./output"
 
INPUT_COL_KODE = "KODE_PROMO"
INPUT_COL_START = "Start"
INPUT_COL_END = "End"
INPUT_COL_JENIS = "JENIS_PROMO"   # opsional, segmen dipisah koma. Lihat parse_jenis_promo()
 
# Nama kolom baru di sheet RAW hasil pemecahan JENIS_PROMO, dan format labelnya.
PROMO_TYPE_COL = "Promo Type"
PROMO_TYPE_LABEL = "{kw} {pct}%"  # {kw} = ADD / DISC, {pct} = angka persen -> "ADD 10%", "DISC 20%"
 
# Kata kategori di JENIS_PROMO -> kode di kolom CATEGORY_COL (sheet RAW). Tambah sesuai kebutuhan.
CATEGORY_COL = "Segment"
CATEGORY_MAP = {
    "SHOES": "SHS", "SHOE": "SHS", "SEPATU": "SHS",
    "BAGS": "BAG", "BAG": "BAG", "HANDBAGS": "BAG", "TAS": "BAG",
}
MD_TOLERANCE = 0.5   # toleransi pembulatan saat membandingkan markdown % (D30 = 29.5-30.5)
INPUT_EXTENSIONS = (".xlsx", ".xlsm", ".xls")
 
# Kalau True, data promo (sheet RAW) juga dibatasi ke tanggal Start-End dari file input.
FILTER_PROMO_QUERY_BY_DATE = True
COUNTER_DEPT_PREFIXES = ("C", "G", "M", "S", "U")
 
DB_DRIVER = "{ODBC Driver 18 for SQL Server}"
DB_SERVER = "149.129.249.131"
DB_DATABASE = "Y2_PROD"
DB_USERNAME = "sa"
DB_PASSWORD = "Eshoes09"
USE_WINDOWS_AUTH = False
 
FONT_NAME = "Arial"
 
def build_query(kode_promo: str, promo_start=None, promo_end=None) -> str:
    kode_sql = str(kode_promo).replace("'", "''")
    date_filter = ""
    if FILTER_PROMO_QUERY_BY_DATE and promo_start is not None and promo_end is not None:
        date_filter = (
            f"AND MLR_DATEPIECE >= '{promo_start:%Y%m%d} 00:00:00' "
            f"AND MLR_DATEPIECE < '{(promo_end + timedelta(days=1)):%Y%m%d} 00:00:00'"
        )
    query = f"""
            SELECT 
            LIGNE.GL_NUMERO as [Invoice Number],
            LIGNE.GL_QTEFACT as [Quantity],
            LIGNE.GL_PUTTC as [Retail Price],
            TYPEREMISE.GTR_LIBELLE as [Disc Description],
            (MLR_TOTALBASE - MLR_MONTANTTTC) as [Selling Price],
            MLR_DATEPIECE as [Document Data],  
            LIGNE.GL_CODEARTICLE as [Item Code],
            MLR_ETABLISSEMENT as [Store of the document], 
            MLR_INDICEG as [Document Index], 
            MLR_MONTANTHT as [Disc - Tax excl. amount], 
            MLR_MONTANTHTDEV as [Disc - Tax excl. amount in currency], 
            MLR_MONTANTTTC as [Disc - Tax incl. amount], 
            MLR_MONTANTTTCDEV as [Disc - Tax incl. amount in currency], 
            MLR_NATUREPIECEG as [Disc line type], 
            MLR_NUMERO as [Doc number], 
            MLR_NUMLIGNE as [Line number], 
            MLR_NUMORDRE as [Order Number], 
            MLR_ORGREMISE as [Origin of discount], 
            MLR_RANG as [Ranking number], 
            MLR_RANGLIE as [Line ranking number], 
            MLR_REMISE as [Disc percentage], 
            MLR_SOUCHE as [Disc line stub], 
            MLR_TOTALBASE as [Disc basis], 
            MLR_TOTALBASEDEV as [Disc basis price - currency], 
            MLR_CODECOND as [Sales Condition],
            LIGNE.GL_ARTICLE as [Item Identiier], 
            LIGNE.GL_CAISSE as [Register of the line], 
            LIGNE.GL_DEPOT as [Warehouse Doc], 
            LIGNE.GL_ETABLISSEMENT as [Store - document line], 
            LIGNE.GL_FAMILLENIV1 as [Brand], 
            LIGNE.GL_FAMILLENIV2 as [Dept], 
            LIGNE.GL_FAMILLENIV3 as [Segment],  
            LIGNE.GL_FOURNISSEUR as [Countermark supplier], 
            LIGNE.GL_LIBELLE as [Line Description], 
            LIGNE.GL_REPRESENTANT as [Doc lin sales representative], 
            LIGNE.GL_TIERS as [Third party of the line], 
            LIGNE.GL_TYPEARTICLE as [Item type line], 
            LIGNE.GL_TYPELIGNE as [Type of line], 
            LIGNE.GL_STATUTENVOI as [Shipping status for items],
            PIECE.GP_CONFIE as [Doc of consigned items],
            ARTICLE.GA_CHARLIBRE1 as [Vendor Style Code], 
            ARTICLE.GA_DESIGNATION1 as [Add item description], 
            ARTICLE.GA_CHARLIBRE2 as [Dimension], 
            ARTICLE.GA_CHARLIBRE3 as [Develop code], 
            GA_CODEBARRE as [Item barcode], 
            GA_CODEDIM1 as [Colour], 
            GA_CODEDIM2 as [Size], 
            GA_CODEDIM3 as [Dimension 3], 
            GA_CODEDIM4 as [Dimension 4], 
            GA_CODEDIM5 as [Dimension 5]
            FROM 
            LIGNEREMISE 
            LEFT JOIN LIGNE ON (
                MLR_NATUREPIECEG = LIGNE.GL_NATUREPIECEG 
                AND MLR_SOUCHE = LIGNE.GL_SOUCHE 
                AND MLR_NUMERO = LIGNE.GL_NUMERO 
                AND MLR_INDICEG = LIGNE.GL_INDICEG 
                AND MLR_NUMLIGNE = LIGNE.GL_NUMLIGNE
            ) 
            LEFT JOIN PIECE ON GP_NATUREPIECEG = LIGNE.GL_NATUREPIECEG 
            AND GP_SOUCHE = LIGNE.GL_SOUCHE 
            AND GP_NUMERO = LIGNE.GL_NUMERO 
            AND GP_INDICEG = LIGNE.GL_INDICEG 
            LEFT JOIN ARTICLE ON GA_ARTICLE = LIGNE.GL_ARTICLE 
            LEFT JOIN ARTICLECOMPL ON GA2_ARTICLE = LIGNE.GL_ARTICLE 
            LEFT JOIN TIERS ON LIGNE.GL_TIERS = T_TIERS 
            AND T_NATUREAUXI = 'CLI'
            LEFT JOIN TYPEREMISE on GTR_TYPEREMISE = MLR_TYPEREMISE
            WHERE 
            MLR_NUMLIGNE > 0 
            AND MLR_NUMORDRE > 0
            AND MLR_CODECOND LIKE '%{kode_sql}%'
            {date_filter}
            ORDER BY MLR_DATEPIECE DESC
    """
    return query
 
def run_query(query: str) -> pd.DataFrame:
    if USE_WINDOWS_AUTH:
        conn_str = (
            f"DRIVER={DB_DRIVER};SERVER={DB_SERVER};DATABASE={DB_DATABASE};"
            f"Trusted_Connection=yes;"
            "Encrypt=yes;"
            "TrustServerCertificate=yes;"
        )
    else:
        conn_str = (
            f"DRIVER={DB_DRIVER};SERVER={DB_SERVER};DATABASE={DB_DATABASE};"
            f"UID={DB_USERNAME};PWD={DB_PASSWORD};"
            "Encrypt=yes;"
            "TrustServerCertificate=yes;"
        )
 
    print("[INFO] Menghubungkan ke database...")
    conn = pyodbc.connect(conn_str)
    try:
        print("[INFO] Menjalankan query...")
        df = pd.read_sql(query, conn)
    finally:
        conn.close()
 
    print(f"[INFO] Query selesai. {len(df)} baris diterima.")
    return df
 
def build_store_sales_query(date_start, date_end, brand: str) -> str:
    if not brand:
        raise ValueError(
            "Brand (GL_FAMILLENIV1) kosong - tidak bisa membangun query Store Sales."
        )
 
    brand_sql = "'" + str(brand).replace("'", "''") + "'"
    date_start_str = f"{date_start:%Y%m%d} 00:00:00"
    date_end_str = f"{(date_end + timedelta(days=1)):%Y%m%d} 00:00:00"
 
    query = f"""
            SELECT 
            GL_ETABLISSEMENT as [Store Code], 
            GL_QTEFACT as [Quantity], 
            GL_TOTALTTC as [Total]
            FROM 
            GCLIGNEARTDIM 
            LEFT OUTER JOIN UTILISAT US1 ON GL_CREATEUR = US1.US_UTILISATEUR 
            WHERE 
            (
                GP_TICKETANNULE <> 'X' 
                AND (
                GP_TYPECOMPTA IN ('TIC')
                ) 
                AND GL_DATEPIECE >= '{date_start_str}' 
                AND GL_DATEPIECE < '{date_end_str}' 
                AND (GL_NATUREPIECEG = 'FFO') 
                AND (
                (
                    GL_TYPELIGNE <> 'COM' 
                    AND GL_TYPELIGNE <> 'TOT'
                )
                ) 
                AND (
                GL_ETABLISSEMENT LIKE '[BCEGHLMRSTUVWZ]%'
                ) 
                AND (
                T_ETABLISSEMENT IN (
                    '10000', 'B010', 'B020', 'B030', 'B040', 'B050', 'B060', 'B070', 'B080', 'B090', 
                    'B100', 'B700', 'C060', 'CWH', 'E010', 'E040', 'E060', 'E070', 'E100', 'E105', 
                    'E120', 'E130', 'E140', 'E141', 'E142', 'E145', 'E150', 'E155', 'E160', 'E170', 
                    'E180', 'E182', 'E185', 'E190', 'E210', 'E220', 'E225', 'E240', 'E250', 'E252', 
                    'E310', 'E315', 'E330', 'E410', 'E415', 'E420', 'E425', 'E450', 'E452', 'E460', 
                    'E510', 'E520', 'E550', 'E552', 'E553', 'E560', 'E570', 'E610', 'E620', 'E630', 
                    'E640', 'E650', 'E660', 'E670', 'E672', 'E680', 'E682', 'E685', 'E690', 'E710', 
                    'E720', 'E730', 'E740', 'E810', 'E815', 'E820', 'E822', 'E850', 'E855', 'E880', 
                    'E885', 'E932', 'E935', 'E940', 'E942', 'E950', 'EDE', 'G035', 'G060', 'H001', 
                    'L100', 'M010', 'M020', 'M030', 'M080', 'M155', 'M190', 'M240', 'M310', 'M420', 
                    'M515', 'M620', 'M730', 'N190', 'N430', 'N930', 'P620', 'R100', 'S010', 'S020', 
                    'S030', 'S040', 'S100', 'S150', 'S185', 'S210', 'S220', 'S230', 'S330', 'S610', 
                    'S630', 'S820', 'S930', 'T045', 'T130', 'T140', 'T150', 'T220', 'T310', 'T340', 
                    'T420', 'T520', 'T560', 'T610', 'T660', 'T685', 'T710', 'T720', 'T730', 'T855', 
                    'T880', 'T950', 'TEST', 'U002', 'U005', 'U195', 'V140', 'V182', 'V210', 'V220', 
                    'V310', 'V415', 'V452', 'V460', 'V520', 'V550', 'V620', 'V630', 'V660', 'V710', 
                    'V720', 'V730', 'V815', 'V820', 'V850', 'V880', 'V935', 'WMKR', 'X999', 'Z010', 
                    'Z030'
                )
                ) 
                AND (T_NATUREAUXI = 'CLI') 
                AND (
                GL_FAMILLENIV1 IN ({brand_sql}) 
                AND GL_FAMILLENIV2 NOT IN ('PKG')
                )
                AND GL_QTEFACT <= 1
            ) 
            ORDER BY 
            GP_DATEPIECE DESC, 
            GL_ETABLISSEMENT, 
            GP_REFINTERNE, 
            GL_NUMLIGNE, 
            GP_HEURECREATION DESC
    """
    return query
 
def aggregate_sales_by_store(sales_df: pd.DataFrame) -> dict:
    """
    Jumlahkan Quantity & Total dari hasil build_store_sales_query(),
    dikelompokkan per Store Code. Hasilnya dict
    {store_code: {"qty": ..., "amount": ...}} yang dipakai untuk mengisi
    kolom Total (periode promo) dan LW Sales per baris toko di SUMMARY.
    """
    if sales_df is None or sales_df.empty:
        return {}
 
    sales_df = sales_df.copy()
    sales_df["Store Code"] = sales_df["Store Code"].astype(str).str.strip()
    sales_df["Quantity"] = pd.to_numeric(sales_df["Quantity"], errors="coerce").fillna(0)
    sales_df["Total"] = pd.to_numeric(sales_df["Total"], errors="coerce").fillna(0)
 
    grouped = sales_df.groupby("Store Code", dropna=True)[["Quantity", "Total"]].sum()
    return {
        store: {"qty": float(row["Quantity"]), "amount": float(row["Total"])}
        for store, row in grouped.iterrows()
    }
 
BULAN_ID = {
    1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "Mei", 6: "Jun",
    7: "Jul", 8: "Agu", 9: "Sep", 10: "Okt", 11: "Nov", 12: "Des",
}
 
HARI_ID = {
    0: "Senin", 1: "Selasa", 2: "Rabu", 3: "Kamis",
    4: "Jumat", 5: "Sabtu", 6: "Minggu",
}
 
# ------------------------------------------------------------
# ATURAN PERIODE LW SALES
# ------------------------------------------------------------
#   1. Jarak dari tanggal akhir LW ke tanggal mulai promo minimal 1 minggu,
#      dihitung INKLUSIF (kedua tanggal ikut dihitung).
#      Contoh: akhir LW 13 Sep -> mulai promo 19 Sep = 13,14,...,19 = 7 hari (OK)
#   2. Selisih periode LW = selisih periode promo (promo_end - promo_start),
#      dan tanggal mulai LW jatuh di hari yang sama dengan mulai promo
#      (mundur per kelipatan 7 hari).
#   3. Periode LW tidak boleh cross dengan tanggal promo
#      (otomatis terpenuhi oleh aturan 1).
#   Dipilih jumlah minggu mundur PALING KECIL yang memenuhi semua aturan.
#
# Contoh:
#   Promo 19 Sep      (selisih 0)  -> LW 12 Sep        (mundur 1 minggu)
#   Promo 19-20 Sep   (selisih 1)  -> LW 12-13 Sep     (mundur 1 minggu)
#   Promo 19-23 Sep   (selisih 4)  -> LW 5-9 Sep       (mundur 2 minggu)
#   Promo 19-25 Sep   (selisih 6)  -> LW 5-11 Sep      (mundur 2 minggu)
#   Promo 19-27 Sep   (selisih 8)  -> LW 5-13 Sep      (mundur 2 minggu)
#   Promo 4-23 Agu    (selisih 19) -> LW 7-26 Jul      (mundur 4 minggu)
LW_MIN_DISTANCE_DAYS = 7  # jarak minimal akhir LW -> mulai promo (inklusif)
 
def lw_distance_days(lw_end, promo_start) -> int:
    """Jarak inklusif dari akhir LW ke mulai promo (13 Sep -> 19 Sep = 7)."""
    return (promo_start - lw_end).days + 1
 
def compute_lw_weeks_back(duration_days: int) -> int:
    """
    Jumlah minggu mundur terkecil (minimal 1) supaya:
        (promo_start - lw_end) + 1 >= LW_MIN_DISTANCE_DAYS
    dengan lw_start = promo_start - 7*k dan lw_end = lw_start + duration_days.
    Dari syarat: 7*k - duration_days + 1 >= LW_MIN_DISTANCE_DAYS
              -> k >= (duration_days + LW_MIN_DISTANCE_DAYS - 1) / 7
    """
    needed = duration_days + LW_MIN_DISTANCE_DAYS - 1
    return max(1, -(-needed // 7))  # ceil tanpa import math
 
def compute_lw_period(promo_start, promo_end):
    """Hitung periode LW dari tanggal Start & End promo (dari file input)."""
    duration_days = (promo_end - promo_start).days  # selisih hari (4-23 Agu = 19)
    weeks_back = compute_lw_weeks_back(duration_days)
 
    lw_start = promo_start - timedelta(weeks=weeks_back)
    lw_end = lw_start + timedelta(days=duration_days)
 
    return lw_start, lw_end
 
def format_date_range_label(prefix: str, start, end) -> str:
    """
    Format label rentang tanggal, mis.:
      - beda bulan : "LW Sales (31 Jul - 7 Agu'25)"
      - bulan sama : "LW Sales (3-9 Jun'26)"
    """
    yy = end.strftime("%y")
 
    if (start.year, start.month) == (end.year, end.month):
        return f"{prefix} ({start.day}-{end.day} {BULAN_ID[end.month]}'{yy})"
    return (
        f"{prefix} ({start.day} {BULAN_ID[start.month]} - "
        f"{end.day} {BULAN_ID[end.month]}'{yy})"
    )
 
def compute_stock_date(promo_start) -> pd.Timestamp:
    """
    Tanggal Stock = 1 hari sebelum tanggal promo dimulai.
    Contoh: promo mulai 7 Agustus 2026 -> Stock 6 Agustus 2026.
    """
    return promo_start - timedelta(days=1)
 
def format_single_date_label(prefix: str, date) -> str:
    """Format label satu tanggal, mis. "Stock 27 Jul 2026"."""
    return f"{prefix} {date.day} {BULAN_ID[date.month]} {date.year}"
 
# ============================================================
# 2. BANGUN TEMPLATE PROMO RESULT (RAW + SUMMARY)
# ============================================================
 
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="FCE4D6")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # kuning muda = sel input manual
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
 
STOCK_HEADER_LABEL = "Stock"
LW_SALES_HEADER_LABEL = "LW Sales"
 
# Kalau True, toko yang punya penjualan brand di periode promo / LW tapi TIDAK
# punya transaksi promo tetap dimunculkan di SUMMARY (kolom diskon = 0).
# Default False: baris toko hanya dari data promo (seperti sebelumnya).
INCLUDE_NON_PROMO_STORES = False
 
def is_counter_dept_store(store_code) -> bool:
    """True kalau kode toko diawali C, G, M, S, atau U (Counter / Dept Store)."""
    return str(store_code).strip().upper().startswith(COUNTER_DEPT_PREFIXES)
 
 
def sort_store_codes(store_codes) -> list:
    """
    Urutkan kode toko: Counter & Dept Store (C, G, M, S, U) dulu,
    baru toko lainnya. Di dalam tiap grup diurutkan alfabetis.
    """
    return sorted(
        store_codes,
        key=lambda st: (0 if is_counter_dept_store(st) else 1, str(st).strip().upper()),
    )
 
 
# ============================================================
# JENIS PROMO -> pecah "Special Offer" jadi label promo (ADD 10%, BUY 2 SP 2,390K, dst.)
# ============================================================
 
def _num(txt: str) -> float:
    """'2,390' / '2.390' / '2390' -> 2390.0 (koma/titik dianggap pemisah ribuan)."""
    return float(re.sub(r"[.,]", "", txt))
 
 
def _md_value(tok: str) -> float:
    """'NML' / 'NORMAL' -> 0, 'D30' -> 30."""
    tok = tok.upper()
    return 0.0 if tok.startswith("N") else float(tok[1:])
 
 
def parse_promo_segment(part: str) -> dict:
    """
    Baca 1 segmen JENIS_PROMO dan ambil semua informasi yang bisa dikenali.
    Informasi yang dipakai untuk MENCOCOKKAN baris data (kriteria):
      pct      : ADD/DISC/EXTRA xx%, atau "B1G1 50%"/"PWP 50%" -> kolom Disc percentage
      md_range : NML-D20, D30-D50, ON D30   -> markdown barang (1 - Disc basis / Retail)
      category : SHOES / BAGS / ...          -> kolom CATEGORY_COL (lihat CATEGORY_MAP)
    Informasi yang hanya dipakai untuk LABEL (tidak bisa dicek per baris):
      BUY n, PAY n GET m, BnGm / BUY n GET m, PWP, SP x,xxxK
    Opsional: "LABEL=teks promo" untuk menentukan label sendiri.
    """
    seg = {"rule": part.strip(), "label": None, "pct": None, "md_range": None,
           "category": None, "mech": [], "has_all_scope": False}
    text = part.strip()
    if "=" in text:
        seg["label"], text = (x.strip() for x in text.split("=", 1))
    up = text.upper()
 
    kw = ""
    m = re.search(r"\b(ADD|ADDITIONAL|EXTRA|DISC(?:OUNT)?)\s*(\d+(?:\.\d+)?)\s*%", up)
    if m:
        kw = "DISC" if m.group(1).startswith("DISC") else "ADD"
        seg["pct"] = float(m.group(2))
 
    m = re.search(r"\b(NML|NORMAL|D\d{1,2})\s*-\s*(NML|D\d{1,2})\b", up)
    if m:
        lo, hi = sorted((_md_value(m.group(1)), _md_value(m.group(2))))
        seg["md_range"] = (lo, hi)
    else:
        m = re.search(r"\bON\s+(?:ALL\s+)?(NML|NORMAL|D\d{1,2})\b", up)
        if m:
            v = _md_value(m.group(1))
            seg["md_range"] = (v, v)
    seg["has_all_scope"] = bool(re.search(r"\bON\s+ALL\b", up))
 
    cats = {code for word, code in CATEGORY_MAP.items() if re.search(rf"\b{word}\b", up)}
    if cats:
        seg["category"] = cats
 
    # --- Mekanisme promo (untuk label & jenis) ---
    m_pwp = re.search(r"\bPWP\b|\bPURCHASE\s+WITH\s+PURCHASE\b", up)
    m_pay = re.search(r"\bPAY\s*(\d+)\s*GET\s*(\d+)", up)
    m_bg = (re.search(r"\bB(\d+)\s*G(\d+)\b", up)
            or re.search(r"\bBUY\s*(\d+)\s*GET\s*(\d+)", up))
    m_buy = re.search(r"\bBUY\s*(\d+)(\s*OR\s*MORE|\s*\+)?", up)
    m_sp = re.search(r"\bSP\.?\s*(\d[\d.,]*)\s*(K|RB|JT)?\b", up)
 
    main = None
    if m_pwp:
        main = "PWP"
    elif m_pay:
        main = f"PAY {m_pay.group(1)} GET {m_pay.group(2)}"
    elif m_bg:
        main = f"B{m_bg.group(1)}G{m_bg.group(2)}"
    if main:
        seg["mech"].append(main)
    elif m_buy:
        seg["mech"].append(f"BUY {m_buy.group(1)}{'+' if m_buy.group(2) else ''}")
    if m_sp:
        seg["mech"].append(f"SP {m_sp.group(1)}{m_sp.group(2) or ''}")
 
    # "B1G1 50%" / "PWP 50% OFF": persen tanpa kata ADD/DISC tetap dibaca untuk promo ini
    if seg["pct"] is None and main:
        m = re.search(r"(\d+(?:\.\d+)?)\s*%", up)
        if m:
            seg["pct"] = float(m.group(1))
 
    seg["kind"] = ("PWP" if m_pwp else "PAY_GET" if m_pay else "BUY_GET" if m_bg
                   else "SPECIAL_PRICE" if m_sp else kw or "LAINNYA")
 
    if not seg["label"]:
        pct_txt = f"{seg['pct']:g}%" if seg["pct"] is not None else ""
        if main and pct_txt:
            seg["label"] = " ".join(x for x in (main, kw, pct_txt) if x)       # "B1G1 50%", "PWP DISC 50%"
        elif main:
            seg["label"] = " ".join(seg["mech"])                              # "PAY 1 GET 2", "PWP SP 99K"
        elif pct_txt:
            seg["label"] = PROMO_TYPE_LABEL.format(kw=kw, pct=f"{seg['pct']:g}")  # "ADD 10%", "DISC 20%"
        elif seg["mech"]:
            seg["label"] = " ".join(seg["mech"])                              # "BUY 2 SP 2,390K"
        else:
            seg["label"] = text[:40]
    return seg
 
 
def _scope_text(seg) -> str:
    parts = []
    if seg["md_range"]:
        lo, hi = seg["md_range"]
        f = lambda v: "NML" if v == 0 else f"D{v:g}"
        parts.append(f(lo) if lo == hi else f"{f(lo)}-{f(hi)}")
    if seg["category"]:
        parts.append("/".join(sorted(seg["category"])))
    return " ".join(parts)
 
 
def parse_jenis_promo(text) -> list:
    """
    Pecah JENIS_PROMO per koma -> list segmen (lihat parse_promo_segment).
    Contoh:
      "BUY2 ADD10%,BUY3 OR MORE ADD15% ON ALL D30-D50" -> ADD 10% | ADD 15%
      "BUY 2 ON D30-D50 SHOES & BAGS SP 2,390K"        -> BUY 2 SP 2,390K
      "PAY 1 GET 2 ON ALL D30-D50.PROMO B1G1 TRACCE"   -> PAY 1 GET 2
      "ADD10% ON NML-D20,ADD15% ON D30-D50"            -> ADD 10% | ADD 15%
      "B1G1 TRACCE=PAY 1 GET 2 ON ALL D30-D50"         -> B1G1 TRACCE (label manual)
    """
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return []
    # Pisah per koma, KECUALI koma pemisah ribuan di antara angka (mis. "SP 2,390K")
    parts = re.split(r"(?<!\d),|,(?!\d)", str(text))
    segs = [parse_promo_segment(p) for p in parts if p.strip()]
 
    # "ON ALL D30-D50" di segmen terakhir berlaku juga untuk segmen lain yang tanpa scope
    last = segs[-1] if segs else None
    if last and last["has_all_scope"]:
        for s in segs[:-1]:
            if s["md_range"] is None and last["md_range"] is not None:
                s["md_range"] = last["md_range"]
            if s["category"] is None and last["category"] is not None:
                s["category"] = last["category"]
 
    # Label kembar -> tambahkan scope, mis. "ADD 10% (NML-D20)"
    labels = [s["label"] for s in segs]
    for s in segs:
        if labels.count(s["label"]) > 1 and _scope_text(s):
            s["label"] = f"{s['label']} ({_scope_text(s)})"
    seen = {}
    for s in segs:
        seen[s["label"]] = seen.get(s["label"], 0) + 1
        if seen[s["label"]] > 1:
            s["label"] = f"{s['label']} #{seen[s['label']]}"
    return segs
 
 
def _segment_mask(df: pd.DataFrame, seg: dict, pct, md):
    """Return (mask baris yang cocok, jumlah kriteria yang dicek)."""
    mask = pd.Series(True, index=df.index)
    n = 0
    if seg["pct"] is not None:
        mask &= (pct - seg["pct"]).abs() < 0.01
        n += 1
    if seg["md_range"] is not None:
        lo, hi = seg["md_range"]
        mask &= md.between(lo - MD_TOLERANCE, hi + MD_TOLERANCE)
        n += 1
    if seg["category"] is not None and CATEGORY_COL in df.columns:
        mask &= df[CATEGORY_COL].astype(str).str.strip().str.upper().isin(seg["category"])
        n += 1
    return mask, n
 
 
def apply_promo_type(df: pd.DataFrame, rules: list):
    """
    Beri label PROMO_TYPE_COL ke setiap baris berdasarkan segmen JENIS_PROMO.
 
    - 1 segmen   : semua baris diberi label segmen tsb (apa pun jenis promonya).
                   Baris yang di luar scope (mis. bukan D30-D50) hanya diberi peringatan.
    - >1 segmen  : tiap baris dicocokkan ke segmen yang SEMUA kriterianya terpenuhi
                   (Disc percentage / markdown / kategori). Kalau cocok ke >1 segmen,
                   dipilih yang kriterianya paling banyak (paling spesifik).
                   Segmen tanpa kriteria yang bisa dicek (mis. SP / PAY 1 GET 2) menjadi
                   penampung baris sisa - asalkan hanya ada 1 segmen seperti itu.
    - Tidak cocok: label = Disc Description asli + persentasenya, dan muncul peringatan.
    Return: (df baru, urutan label untuk kolom SUMMARY)
    """
    df = df.copy()
    pct = pd.to_numeric(df.get("Disc percentage"), errors="coerce")
    qty = pd.to_numeric(df["Quantity"], errors="coerce").replace(0, pd.NA)
    retail = pd.to_numeric(df["Retail Price"], errors="coerce").replace(0, pd.NA)
    basis = (pd.to_numeric(df["Disc basis"], errors="coerce") / qty).abs()
    md = ((1 - basis / retail) * 100).astype(float).round(1)
    df["Markdown %"] = md
 
    fallback = df["Disc Description"].astype(str) + pct.map(
        lambda v: "" if pd.isna(v) else f" {v:g}%")
    label = pd.Series(pd.NA, index=df.index, dtype=object)
 
    if len(rules) == 1:
        seg = rules[0]
        label[:] = seg["label"]
        mask, n = _segment_mask(df, seg, pct, md)
        if n and (~mask).any():
            print(f"[WARNING] {int((~mask).sum())} baris tidak sesuai scope '{seg['rule']}' "
                  f"(cek kolom Markdown % / Disc percentage) - tetap dilabeli '{seg['label']}'.")
    else:
        best_score = pd.Series(-1, index=df.index)
        catch_all = [s for s in rules if _segment_mask(df, s, pct, md)[1] == 0]
        for seg in rules:
            mask, n = _segment_mask(df, seg, pct, md)
            if n == 0:
                continue
            take = mask & (n > best_score)
            label[take] = seg["label"]
            best_score[take] = n
        if len(catch_all) == 1:
            label = label.fillna(catch_all[0]["label"])
        elif len(catch_all) > 1:
            print(
                f"[WARNING] {len(catch_all)} segmen JENIS_PROMO tidak punya kriteria yang bisa "
                f"dicocokkan ke data ({[s['label'] for s in catch_all]}). Tambahkan ADDxx% / "
                f"Dxx-Dyy / kategori, atau data tidak bisa dipecah otomatis."
            )
 
    n_unmatched = int(label.isna().sum())
    df[PROMO_TYPE_COL] = label.fillna(fallback)
    if n_unmatched:
        print(f"[WARNING] {n_unmatched} baris tidak cocok dengan JENIS_PROMO - "
              f"dilabeli dengan Disc Description aslinya.")
    for s in rules:
        n = int((df[PROMO_TYPE_COL] == s["label"]).sum())
        print(f"[INFO]   {s['label']:<18} <- '{s['rule']}' : {n} baris")
 
    order = [s["label"] for s in rules]
    order += sorted(set(df[PROMO_TYPE_COL]) - set(order))
    return df, order
 
 
def build_promo_result_workbook(
    df: pd.DataFrame,
    output_path: str,
    lw_sales_label: str = None,
    stock_label: str = None,
    lw_sales_map: dict = None,
    total_sales_label: str = None,
    total_sales_map: dict = None,
    group_col: str = "Disc Description",
    group_order: list = None,
    jenis_promo_text: str = None,
):
    # group_col  : kolom RAW untuk memecah kolom diskon di SUMMARY
    #              ("Disc Description" default, atau PROMO_TYPE_COL kalau JENIS_PROMO diisi).
    # group_order: urutan kolom diskon (default: alfabetis).
    # Kalau label tidak dikirim (mis. dipanggil manual tanpa lewat main()),
    # pakai fallback konstanta di bagian KONFIGURASI.
    lw_sales_label = lw_sales_label or LW_SALES_HEADER_LABEL
    stock_label = stock_label or STOCK_HEADER_LABEL
    total_sales_label = total_sales_label or "Total"
    lw_sales_map = lw_sales_map or {}
    total_sales_map = total_sales_map or {}
 
    required_cols = ["Store of the document", group_col, "Quantity", "Selling Price"]
    for col in required_cols:
        if col not in df.columns:
            raise KeyError(f"Kolom '{col}' tidak ada di hasil query.")
 
    df = df.dropna(subset=["Store of the document", group_col]).copy()
    df["Quantity"] = df["Quantity"].fillna(0)
    df["Selling Price"] = df["Selling Price"].fillna(0)
 
    present = set(df[group_col].unique().tolist())
    if group_order:
        disc_set = [g for g in group_order if g in present]
        disc_set += sorted(present - set(disc_set))
    else:
        disc_set = sorted(present)
    store_set = sort_store_codes(df["Store of the document"].unique().tolist())
 
    # Toko yang punya penjualan brand (periode promo / LW) tapi tidak ada di data promo
    promo_stores_clean = {str(st).strip() for st in store_set}
    extra_stores = sorted(
        (set(total_sales_map) | set(lw_sales_map)) - promo_stores_clean
    )
    if extra_stores:
        if INCLUDE_NON_PROMO_STORES:
            print(f"[INFO] {len(extra_stores)} toko tanpa transaksi promo ikut ditampilkan: {extra_stores}")
            store_set = sort_store_codes(set(store_set) | set(extra_stores))
        else:
            print(
                f"[WARNING] {len(extra_stores)} toko punya penjualan brand tapi tidak ada di data "
                f"promo, tidak ditampilkan di SUMMARY."
                # f" {extra_stores} (set INCLUDE_NON_PROMO_STORES = True untuk menampilkannya)."
            )
 
    out_wb = openpyxl.Workbook()
 
    raw_ws = out_wb.active
    raw_ws.title = "RAW"
    all_cols = list(df.columns)
    raw_ws.append(all_cols)
    for cell in raw_ws[1]:
        cell.font = Font(name=FONT_NAME, bold=True)
    for row in df.itertuples(index=False, name=None):
        raw_ws.append(list(row))
    for c_idx in range(1, len(all_cols) + 1):
        raw_ws.column_dimensions[get_column_letter(c_idx)].width = 20
    raw_ws.freeze_panes = "A2"
 
    # Cari posisi kolom (di sheet RAW) yang dibutuhkan formula SUMMARY,
    # supaya tetap benar walau urutan/isi kolom RAW berubah.
    store_col_letter = get_column_letter(all_cols.index("Store of the document") + 1)
    disc_col_letter = get_column_letter(all_cols.index(group_col) + 1)
    qty_col_letter = get_column_letter(all_cols.index("Quantity") + 1)
    amt_col_letter = get_column_letter(all_cols.index("Selling Price") + 1)
 
    raw_rows = len(df)
    RAW_STORE_RNG = f"RAW!${store_col_letter}$2:${store_col_letter}${raw_rows + 1}"
    RAW_DISC_RNG = f"RAW!${disc_col_letter}$2:${disc_col_letter}${raw_rows + 1}"
    RAW_QTY_RNG = f"RAW!${qty_col_letter}$2:${qty_col_letter}${raw_rows + 1}"
    RAW_AMT_RNG = f"RAW!${amt_col_letter}$2:${amt_col_letter}${raw_rows + 1}"
 
    # --- SUMMARY sheet ---
    sum_ws = out_wb.create_sheet("SUMMARY")
 
    n_disc = len(disc_set)
    # total_cols awal hanya dipakai untuk menulis judul; akan disesuaikan lagi
    # (di-merge ulang) setelah kolom Stock/Sellthru/Cont%/LW Sales/Growth% dibuat.
    sum_ws.cell(row=1, column=1, value="PROMO RESULT").font = Font(name=FONT_NAME, bold=True, size=14)
 
    note = sum_ws.cell(
        row=2, column=1,
        value=(
            f"Kolom diskon: Qty = SUM(Quantity), Amount = SUM(Selling Price) per Store x "
            f"{group_col} (sheet RAW). Total & LW Sales: penjualan brand per toko selama periode "
            "promo & periode LW (query sales toko)."
        ),
    )
    note.font = Font(name=FONT_NAME, italic=True, size=9, color="808080")
    if jenis_promo_text:
        jp = sum_ws.cell(row=3, column=1, value=f"Jenis Promo: {jenis_promo_text}")
        jp.font = Font(name=FONT_NAME, bold=True, size=10)
 
    HEADER_ROW = 4
    SUBHEADER_ROW = 5
    DATA_START_ROW = 6
 
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=1, end_row=SUBHEADER_ROW, end_column=1)
    c = sum_ws.cell(row=HEADER_ROW, column=1, value="Store Code")
    c.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER
 
    col = 2
    disc_col_start = {}
    for disc in disc_set:
        sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
        hc = sum_ws.cell(row=HEADER_ROW, column=col, value=disc)
        hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        hc.fill = HEADER_FILL
        hc.alignment = Alignment(horizontal="center", vertical="center")
        for cc in range(col, col + 2):
            sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
 
        qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=col, value="Qty")
        amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=col + 1, value="Amount")
        for cell in (qty_c, amt_c):
            cell.font = Font(name=FONT_NAME, bold=True)
            cell.fill = SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
 
        disc_col_start[disc] = col
        col += 2
 
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
    hc = sum_ws.cell(row=HEADER_ROW, column=col, value=total_sales_label)
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cc in range(col, col + 2):
        sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
    total_qty_col = col
    total_amt_col = col + 1
    qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=total_qty_col, value="Qty")
    amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=total_amt_col, value="Amount")
    for cell in (qty_c, amt_c):
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
 
    # --- Kolom tambahan: Stock | Sellthru | Cont% | LW Sales (Qty/Amount) | Growth% ---
    col = total_amt_col + 1
 
    stock_col = col
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=stock_col, end_row=SUBHEADER_ROW, end_column=stock_col)
    hc = sum_ws.cell(row=HEADER_ROW, column=stock_col, value=stock_label)
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr in (HEADER_ROW, SUBHEADER_ROW):
        sum_ws.cell(row=rr, column=stock_col).border = BORDER
    col += 1
 
    sellthru_col = col
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=sellthru_col, end_row=SUBHEADER_ROW, end_column=sellthru_col)
    hc = sum_ws.cell(row=HEADER_ROW, column=sellthru_col, value="Sellthru")
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr in (HEADER_ROW, SUBHEADER_ROW):
        sum_ws.cell(row=rr, column=sellthru_col).border = BORDER
    col += 1
 
    cont_col = col
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=cont_col, end_row=SUBHEADER_ROW, end_column=cont_col)
    hc = sum_ws.cell(row=HEADER_ROW, column=cont_col, value="Cont%")
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr in (HEADER_ROW, SUBHEADER_ROW):
        sum_ws.cell(row=rr, column=cont_col).border = BORDER
    col += 1
 
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=col, end_row=HEADER_ROW, end_column=col + 1)
    hc = sum_ws.cell(row=HEADER_ROW, column=col, value=lw_sales_label)
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cc in range(col, col + 2):
        sum_ws.cell(row=HEADER_ROW, column=cc).border = BORDER
    lw_qty_col = col
    lw_amt_col = col + 1
    qty_c = sum_ws.cell(row=SUBHEADER_ROW, column=lw_qty_col, value="Qty")
    amt_c = sum_ws.cell(row=SUBHEADER_ROW, column=lw_amt_col, value="Amount")
    for cell in (qty_c, amt_c):
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    col += 2
 
    growth_col = col
    sum_ws.merge_cells(start_row=HEADER_ROW, start_column=growth_col, end_row=SUBHEADER_ROW, end_column=growth_col)
    hc = sum_ws.cell(row=HEADER_ROW, column=growth_col, value="Growth%")
    hc.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    hc.fill = HEADER_FILL
    hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr in (HEADER_ROW, SUBHEADER_ROW):
        sum_ws.cell(row=rr, column=growth_col).border = BORDER
 
    total_cols = growth_col  # update total_cols agar merge title/note di atas ikut menyesuaikan
    sum_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    sum_ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
 
    # baris grand total sudah diketahui sebelum loop, dipakai untuk formula Cont%
    grand_row = DATA_START_ROW + len(store_set)
 
    r = DATA_START_ROW
    for store in store_set:
        sc = sum_ws.cell(row=r, column=1, value=store)
        sc.font = Font(name=FONT_NAME)
        sc.alignment = Alignment(horizontal="center")
        sc.border = BORDER
 
        for disc in disc_set:
            qcol = disc_col_start[disc]
            acol = qcol + 1
            qty_formula = (
                f'=SUMIFS({RAW_QTY_RNG},{RAW_STORE_RNG},$A{r},'
                f'{RAW_DISC_RNG},{get_column_letter(qcol)}$4)'
            )
            amt_formula = (
                f'=SUMIFS({RAW_AMT_RNG},{RAW_STORE_RNG},$A{r},'
                f'{RAW_DISC_RNG},{get_column_letter(qcol)}$4)'
            )
            qc = sum_ws.cell(row=r, column=qcol, value=qty_formula)
            ac = sum_ws.cell(row=r, column=acol, value=amt_formula)
            qc.number_format = '#,##0'
            ac.number_format = '#,##0'
            qc.font = Font(name=FONT_NAME)
            ac.font = Font(name=FONT_NAME)
            qc.border = BORDER
            ac.border = BORDER
 
        # --- Total Qty/Amount: penjualan brand toko ini selama periode promo,
        #     dari query store sales (lihat aggregate_sales_by_store()).
        #     Kalau toko tidak ketemu di hasil query, dianggap 0.
        total_data = total_sales_map.get(str(store).strip(), {})
        tq = sum_ws.cell(row=r, column=total_qty_col, value=total_data.get("qty", 0))
        ta = sum_ws.cell(row=r, column=total_amt_col, value=total_data.get("amount", 0))
        tq.number_format = '#,##0'
        ta.number_format = '#,##0'
        tq.font = Font(name=FONT_NAME, bold=True)
        ta.font = Font(name=FONT_NAME, bold=True)
        tq.border = BORDER
        ta.border = BORDER
 
        # --- Stock: sel input manual (kosong, isi nanti kalau sumber datanya sudah ada) ---
        stock_c = sum_ws.cell(row=r, column=stock_col, value=None)
        stock_c.fill = INPUT_FILL
        stock_c.border = BORDER
        stock_c.number_format = '#,##0'
 
        # --- Sellthru = Total Qty / Stock (akan #DIV/0! selama Stock masih kosong) ---
        tqty_ref = f"{get_column_letter(total_qty_col)}{r}"
        stock_ref = f"{get_column_letter(stock_col)}{r}"
        sellthru_c = sum_ws.cell(row=r, column=sellthru_col, value=f"={tqty_ref}/{stock_ref}")
        sellthru_c.number_format = '0.0%'
        sellthru_c.font = Font(name=FONT_NAME)
        sellthru_c.border = BORDER
 
        # --- Cont% = Total Qty toko ini / Grand Total Qty (lihat catatan asumsi di atas) ---
        grand_qty_ref = f"${get_column_letter(total_qty_col)}${grand_row}"
        cont_c = sum_ws.cell(row=r, column=cont_col, value=f"={tqty_ref}/{grand_qty_ref}")
        cont_c.number_format = '0.0%'
        cont_c.font = Font(name=FONT_NAME)
        cont_c.border = BORDER
 
        # --- LW Sales Qty/Amount: sekarang otomatis dari query LW Sales,
        #     dijumlahkan per Store Code (lihat aggregate_sales_by_store()).
        #     Kalau toko tidak ketemu di hasil query LW, dianggap 0.
        lw_data = lw_sales_map.get(str(store).strip(), {})
        lwq_c = sum_ws.cell(row=r, column=lw_qty_col, value=lw_data.get("qty", 0))
        lwa_c = sum_ws.cell(row=r, column=lw_amt_col, value=lw_data.get("amount", 0))
        for cell in (lwq_c, lwa_c):
            cell.font = Font(name=FONT_NAME)
            cell.border = BORDER
            cell.number_format = '#,##0'
 
        # --- Growth% = Total Amount / LW Amount - 1 (akan #DIV/0! selama LW Amount kosong) ---
        tamt_ref = f"{get_column_letter(total_amt_col)}{r}"
        lwamt_ref = f"{get_column_letter(lw_amt_col)}{r}"
        growth_c = sum_ws.cell(row=r, column=growth_col, value=f"={tamt_ref}/{lwamt_ref}-1")
        growth_c.number_format = '0.0%'
        growth_c.font = Font(name=FONT_NAME)
        growth_c.border = BORDER
 
        r += 1
 
    last_data_row = r - 1
    assert grand_row == r, "grand_row seharusnya sama dengan baris setelah loop store"
 
    gt = sum_ws.cell(row=r, column=1, value="GRAND TOTAL")
    gt.font = Font(name=FONT_NAME, bold=True)
    gt.fill = TOTAL_FILL
    gt.border = BORDER
    for cc in range(2, total_qty_col):  # sebelum kolom Total: SUM per kolom Qty/Amount disc
        col_letter = get_column_letter(cc)
        cell = sum_ws.cell(row=r, column=cc, value=f"=SUM({col_letter}{DATA_START_ROW}:{col_letter}{last_data_row})")
        cell.number_format = '#,##0'
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    for cc in (total_qty_col, total_amt_col, stock_col, lw_qty_col, lw_amt_col):
        col_letter = get_column_letter(cc)
        cell = sum_ws.cell(row=r, column=cc, value=f"=SUM({col_letter}{DATA_START_ROW}:{col_letter}{last_data_row})")
        cell.number_format = '#,##0'
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    # Sellthru / Cont% / Growth% di baris Grand Total dihitung ulang dari total-nya (bukan SUM dari %)
    gt_tqty_ref = f"{get_column_letter(total_qty_col)}{r}"
    gt_tamt_ref = f"{get_column_letter(total_amt_col)}{r}"
    gt_stock_ref = f"{get_column_letter(stock_col)}{r}"
    gt_lwamt_ref = f"{get_column_letter(lw_amt_col)}{r}"
    for cc, formula in (
        (sellthru_col, f"={gt_tqty_ref}/{gt_stock_ref}"),
        (cont_col, "=1"),
        (growth_col, f"={gt_tamt_ref}/{gt_lwamt_ref}-1"),
    ):
        cell = sum_ws.cell(row=r, column=cc, value=formula)
        cell.number_format = '0.0%'
        cell.font = Font(name=FONT_NAME, bold=True)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
        cell.border = BORDER
 
    sum_ws.column_dimensions["A"].width = 14
    for cc in range(2, total_cols + 1):
        sum_ws.column_dimensions[get_column_letter(cc)].width = 16
    sum_ws.freeze_panes = "B6"
 
    out_wb.save(output_path)
    print(f"[INFO] File tersimpan: {output_path}")
    print(f"[INFO] Disc types: {disc_set}")
    print(f"[INFO] Jumlah store: {len(store_set)}")
 
# ============================================================
# 3. BACA FILE INPUT
# ============================================================
 
def _norm(text) -> str:
    return str(text).strip().lower().replace(" ", "").replace("_", "")
 
 
def _clean_kode_promo(value):
    """6988 / 6988.0 / ' 6988 ' -> '6988'. Kosong -> None."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    kode = str(value).strip()
    if kode.endswith(".0") and kode[:-2].isdigit():
        kode = kode[:-2]
    return kode or None
 
 
def _parse_date(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value).normalize()
    ts = pd.to_datetime(str(value).strip(), dayfirst=True, errors="coerce")
    return None if pd.isna(ts) else ts.normalize()
 
 
def read_input_file(path: str) -> list:
    """
    Baca 1 file input. Baris 1 = header (KODE_PROMO | Start | End | JENIS_PROMO opsional),
    data promo dimulai dari baris setelah header.
    Return: list of dict {kode, start, end, source}.
    """
    raw = pd.read_excel(path, sheet_name=0, dtype=object)
    if raw.empty:
        return []
 
    cols = {_norm(c): c for c in raw.columns}
    try:
        col_kode = cols[_norm(INPUT_COL_KODE)]
        col_start = cols[_norm(INPUT_COL_START)]
        col_end = cols[_norm(INPUT_COL_END)]
    except KeyError:
        if raw.shape[1] < 3:
            raise ValueError(
                f"File {os.path.basename(path)} harus punya minimal 3 kolom "
                f"({INPUT_COL_KODE}, {INPUT_COL_START}, {INPUT_COL_END})."
            )
        print(
            f"[WARNING] Header '{INPUT_COL_KODE}/{INPUT_COL_START}/{INPUT_COL_END}' tidak "
            f"lengkap di {os.path.basename(path)} - dipakai kolom A, B, C."
        )
        col_kode, col_start, col_end = raw.columns[:3]
 
    col_jenis = cols.get(_norm(INPUT_COL_JENIS))  # kolom opsional
 
    promos = []
    for idx, row in raw.iterrows():
        excel_row = idx + 2  # +1 header, +1 karena index mulai 0
        kode = _clean_kode_promo(row[col_kode])
        start = _parse_date(row[col_start])
        end = _parse_date(row[col_end])
 
        if kode is None and start is None and end is None:
            continue  # baris kosong
        if kode is None or start is None or end is None:
            print(
                f"[WARNING] {os.path.basename(path)} baris {excel_row}: data tidak lengkap / "
                f"tanggal tidak valid (kode={row[col_kode]}, start={row[col_start]}, "
                f"end={row[col_end]}). Dilewati."
            )
            continue
        if end < start:
            print(
                f"[WARNING] {os.path.basename(path)} baris {excel_row}: tanggal End "
                f"({end:%d %b %Y}) sebelum Start ({start:%d %b %Y}). Dilewati."
            )
            continue
 
        jenis = None
        if col_jenis is not None and not pd.isna(row[col_jenis]):
            jenis = str(row[col_jenis]).strip() or None
 
        promos.append({"kode": kode, "start": start, "end": end, "jenis": jenis,
                       "source": f"{os.path.basename(path)} baris {excel_row}"})
    return promos
 
 
def read_all_inputs(folder: str) -> list:
    if not os.path.isdir(folder):
        raise FileNotFoundError(f"Folder input tidak ditemukan: {folder}")
 
    files = sorted(
        f for f in os.listdir(folder)
        if f.lower().endswith(INPUT_EXTENSIONS) and not f.startswith("~$")  # skip file lock Excel
    )
    if not files:
        print(f"[WARNING] Tidak ada file Excel di folder input: {folder}")
        return []
 
    promos = []
    for f in files:
        path = os.path.join(folder, f)
        print(f"[INFO] Membaca file input: {path}")
        try:
            found = read_input_file(path)
        except Exception as e:
            print(f"[ERROR] Gagal membaca {f}: {e}")
            continue
        print(f"[INFO]   -> {len(found)} promo ditemukan.")
        promos.extend(found)
    return promos
 
 
# ============================================================
# 4. PROSES 1 PROMO
# ============================================================
 
def process_promo(
    kode_promo: str,
    promo_start,
    promo_end,
    jenis_promo: str = None,
    output_dir: str = None,
):
    print("=" * 70)
    print(f"[INFO] KODE PROMO: {kode_promo} | {promo_start:%d %b %Y} - {promo_end:%d %b %Y}")
    print("=" * 70)
 
    # 1. Data promo (sheet RAW)
    df = run_query(build_query(kode_promo, promo_start, promo_end))
    if df.empty:
        print(f"[WARNING] Tidak ada data promo untuk kode {kode_promo}. Dilewati.")
        return None
 
    # 1b. Pecah diskon berdasarkan JENIS_PROMO (kalau diisi di file input)
    group_col, group_order = "Disc Description", None
    rules = parse_jenis_promo(jenis_promo)
    if rules:
        print(f"[INFO] JENIS_PROMO: {jenis_promo}")
        df, group_order = apply_promo_type(df, rules)
        group_col = PROMO_TYPE_COL
    elif jenis_promo:
        print("[WARNING] JENIS_PROMO diisi tapi kosong setelah dibaca - pakai Disc Description.")
 
    # 2. Periode Total / LW / Stock - semuanya dari tanggal di file input
    duration_days = (promo_end - promo_start).days
    lw_start, lw_end = compute_lw_period(promo_start, promo_end)
    lw_sales_label = format_date_range_label("LW Sales", lw_start, lw_end)
    total_sales_label = format_date_range_label("Total", promo_start, promo_end)
 
    stock_date = compute_stock_date(promo_start)
    stock_label = format_single_date_label("Stock", stock_date)
 
    print(
        f"[INFO] Periode promo (input)   : {promo_start:%d %b %Y} - {promo_end:%d %b %Y} "
        f"({duration_days} hari, mulai hari {HARI_ID[promo_start.weekday()]})"
    )
    print(
        f"[INFO] Periode LW Sales        : {lw_start:%d %b %Y} - {lw_end:%d %b %Y} "
        f"(mundur {compute_lw_weeks_back(duration_days)} minggu, "
        f"jarak {lw_distance_days(lw_end, promo_start)} hari ke mulai promo)"
    )
    print(f"[INFO] Tanggal Stock (H-1 promo): {stock_date:%d %b %Y}")
 
    # 3. Sales toko periode promo & LW
    brands_found = sorted(df["Brand"].dropna().astype(str).str.strip().unique().tolist())
    if not brands_found:
        print("[WARNING] Kolom 'Brand' kosong di data promo - Total & LW Sales dianggap 0 semua.")
        total_sales_map, lw_sales_map = {}, {}
    else:
        if len(brands_found) > 1:
            print(
                f"[WARNING] Ditemukan {len(brands_found)} Brand berbeda: {brands_found}. "
                f"Dipakai brand pertama saja: '{brands_found[0]}'."
            )
        brand = brands_found[0]
        print(f"[INFO] Brand untuk query Store Sales (GL_FAMILLENIV1): '{brand}'")
 
        print("[INFO] Query sales periode PROMO (kolom Total)...")
        total_df = run_query(build_store_sales_query(promo_start, promo_end, brand))
        total_sales_map = aggregate_sales_by_store(total_df)
        print(f"[INFO] Data sales periode promo: {len(total_df)} baris, {len(total_sales_map)} store.")
 
        print("[INFO] Query sales periode LW (kolom LW Sales)...")
        lw_df = run_query(build_store_sales_query(lw_start, lw_end, brand))
        lw_sales_map = aggregate_sales_by_store(lw_df)
        print(f"[INFO] Data sales periode LW: {len(lw_df)} baris, {len(lw_sales_map)} store.")
 
 
    output_filename = (
        f"Promo_Result_{kode_promo}_{datetime.now():%d_%b_%Y}.xlsx"
    )

    target_output_folder = output_dir or OUTPUT_FOLDER
    os.makedirs(target_output_folder, exist_ok=True)

    output_path = os.path.join(
        target_output_folder,
        output_filename,
    )
     
       
    build_promo_result_workbook(
        df,
        output_path,
        lw_sales_label=lw_sales_label,
        stock_label=stock_label,
        lw_sales_map=lw_sales_map,
        total_sales_label=total_sales_label,
        total_sales_map=total_sales_map,
        group_col=group_col,
        group_order=group_order,
        jenis_promo_text=jenis_promo,
    )
    return output_path
 
 
def main():
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
 
    promos = read_all_inputs(INPUT_FOLDER)
    if not promos:
        print("[WARNING] Tidak ada promo valid di file input. Proses dihentikan.")
        sys.exit(0)
 
    print(f"[INFO] Total promo yang akan diproses: {len(promos)}")
 
    success, failed = [], []
    for p in promos:
        try:
            out = process_promo(p["kode"], p["start"], p["end"], p.get("jenis"))
            if out:
                success.append((p["kode"], out))
            else:
                failed.append((p["kode"], "tidak ada data promo"))
        except Exception as e:
            print(f"[ERROR] Kode promo {p['kode']} ({p['source']}) gagal: {e}")
            failed.append((p["kode"], str(e)))
 
    print("=" * 70)
    print(f"[INFO] SELESAI. Berhasil: {len(success)} | Gagal/dilewati: {len(failed)}")
    for kode, out in success:
        print(f"   [OK]   {kode} -> {out}")
    for kode, reason in failed:
        print(f"   [SKIP] {kode} -> {reason}")

 
def run_promo_result(input_path: str, output_dir: str) -> list[str]:
    """Entry point yang dipanggil oleh Flask."""

    if not os.path.isfile(input_path):
        raise FileNotFoundError(
            f"File input tidak ditemukan: {input_path}"
        )

    os.makedirs(output_dir, exist_ok=True)

    promos = read_input_file(input_path)

    if not promos:
        raise ValueError(
            "Tidak ada baris promo valid pada file Excel."
        )

    results = []

    for promo in promos:
        try:
            output_path = process_promo(
                kode_promo=promo["kode"],
                promo_start=promo["start"],
                promo_end=promo["end"],
                jenis_promo=promo.get("jenis"),
                output_dir=output_dir,
            )

            if output_path and os.path.isfile(output_path):
                results.append(output_path)

        except Exception as exc:
            raise RuntimeError(
                f"Gagal memproses kode promo "
                f"{promo.get('kode', '(tidak diketahui)')}: {exc}"
            ) from exc

    if not results:
        raise RuntimeError(
            "Tidak ada file Promo Result yang berhasil dibuat. "
            "Periksa kode promo dan hasil query database."
        )

    return results
 
if __name__ == "__main__":
    main()